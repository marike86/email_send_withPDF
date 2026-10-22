"""PDF 批量邮件的共用业务逻辑和命令行入口。"""

import argparse
import csv
import getpass
import hashlib
import json
import math
import os
from pathlib import Path
import re
import smtplib
import ssl
import tempfile
import threading
from zipfile import BadZipFile
from contextlib import contextmanager
from datetime import datetime
from email.message import EmailMessage

import openpyxl

DEFAULT_CONFIG = {
    "smtp_server": "smtp.gmail.com", "smtp_port": 587, "use_tls": True,
    "use_ssl": False, "email_template": "email_template.txt",
    "student_info_excel": "student_info.xlsx", "send_interval": 3,
    "max_retries": 3, "retry_interval": 10,
    "excel_student_id_col": 1, "excel_name_col": 2, "excel_email_col": 3,
}
STATUS_FILE = "status.json"
LABELS = {"ready": "待发送", "success": "已发送", "failed": "失败",
          "blocked": "需修正", "sending": "待核实", "uncertain": "待核实"}


def clean_path(path):
    """清理用户粘贴的路径，例如带引号的 Windows 路径。"""
    return str(path or "").strip().strip("\"'").strip()


def atomic_json(path, data):
    """同目录写入临时文件后替换，避免中断造成 JSON 截断。"""
    target = Path(path).resolve()
    target.parent.mkdir(parents=True, exist_ok=True)
    fd, name = tempfile.mkstemp(dir=target.parent, prefix=target.name, suffix=".tmp")
    try:
        with os.fdopen(fd, "w", encoding="utf-8") as stream:
            json.dump(data, stream, ensure_ascii=False, indent=2)
            stream.flush()
            os.fsync(stream.fileno())
        os.replace(name, target)
    finally:
        if os.path.exists(name):
            os.unlink(name)


def load_config(config_path="config.json"):
    """加载配置；相对文件路径以配置文件目录为基准，环境变量优先。"""
    path = Path(config_path).resolve()
    config = DEFAULT_CONFIG.copy()
    if path.exists():
        with path.open(encoding="utf-8-sig") as stream:
            loaded = json.load(stream)
        if not isinstance(loaded, dict):
            raise ValueError("配置文件必须是 JSON 对象")
        config.update(loaded)
    for key in DEFAULT_CONFIG.keys() | {"sender_email", "sender_password", "attachment_folder"}:
        value = os.getenv(key.upper())
        if value is None and key.startswith("smtp_"):
            value = os.getenv("SMTP_" + key.upper())  # 兼容旧变量名
        if value is not None:
            if key in {"use_tls", "use_ssl"}:
                if value.lower() not in {"true", "false", "1", "0"}:
                    raise ValueError(f"{key.upper()} 应为 true/false")
                value = value.lower() in {"true", "1"}
            elif key in {"excel_student_id_col", "excel_name_col", "excel_email_col"}:
                value = column_number(value)
            elif key in {"smtp_port", "max_retries"}:
                value = int(value)
            elif key in {"send_interval", "retry_interval"}:
                value = float(value)
            config[key] = value
    for key in ("email_template", "student_info_excel", "attachment_folder"):
        if config.get(key):
            value = Path(clean_path(config[key])).expanduser()
            config[key] = str(value if value.is_absolute() else path.parent / value)
    return config


def save_config(config, config_path="config.json"):
    """保存非敏感配置；授权码仅保留在内存或环境变量中。"""
    atomic_json(config_path, {k: v for k, v in config.items() if k != "sender_password"})


def column_number(value):
    """将 A、AA 或整数转换为有效 Excel 列号。"""
    text = str(value).strip().upper()
    number = int(text) if text.isdigit() else openpyxl.utils.column_index_from_string(text)
    if not 1 <= number <= 16384:
        raise ValueError("Excel 列应为 A 至 XFD 或 1 至 16384")
    return number


def valid_email(value):
    """仅接受单个邮箱地址，拒绝换行、列表和显示名称。"""
    return bool(re.fullmatch(r"[^\s@<>,;]+@[^\s@<>,;]+\.[^\s@<>,;]+", str(value)))


def validate_config(config, sending=False):
    """校验发送节奏及 SMTP 参数；预检无需授权码。"""
    for key in ("send_interval", "retry_interval"):
        number = float(config.get(key, DEFAULT_CONFIG[key]))
        if not math.isfinite(number) or number < 0:
            raise ValueError(f"{key} 必须是非负有限数")
    if not 1 <= int(config.get("max_retries", 3)) <= 10:
        raise ValueError("尝试次数必须在 1 至 10 之间")
    if not 1 <= int(config["smtp_port"]) <= 65535:
        raise ValueError("SMTP 端口必须在 1 至 65535 之间")
    if not str(config.get("smtp_server", "")).strip():
        raise ValueError("请输入 SMTP 服务器")
    if sending:
        if not valid_email(config.get("sender_email", "")):
            raise ValueError("请输入有效的发件人邮箱")
        if not config.get("sender_password"):
            raise ValueError("请输入邮箱授权码")
        if not (config.get("use_tls") or config.get("use_ssl") or int(config["smtp_port"]) == 465):
            raise ValueError("请启用 STARTTLS 或 SSL，以保护授权码")


def extract_student_id(filename):
    """提取 2023001_001.pdf 或 签署后：2023001_001.pdf 中的学号。"""
    match = re.fullmatch(r"(?:签署后[：:]\s*)?(\d+)(?:_.+)?", Path(filename).stem)
    return match.group(1) if match else None


def get_files_with_numbers(folder_path):
    """列出 PDF；未匹配学号的 PDF 也返回以便预检报告。"""
    folder = Path(clean_path(folder_path)).resolve()
    if not folder.is_dir():
        raise FileNotFoundError(f"附件文件夹不存在: {folder}")
    return [{"filename": p.name, "filepath": str(p), "student_id": extract_student_id(p.name)}
            for p in sorted(folder.iterdir()) if p.is_file() and p.suffix.lower() == ".pdf"]


def load_students(excel_path, student_id_col=1, name_col=2, email_col=3):
    """一次读取并关闭 Excel；拒绝重复学号，保留文本学号的前导零。"""
    columns = [column_number(x) for x in (student_id_col, name_col, email_col)]
    if len(set(columns)) != 3:
        raise ValueError("学号、姓名、邮箱必须使用不同的列")
    try:
        wb = openpyxl.load_workbook(clean_path(excel_path), read_only=True, data_only=True)
    except (OSError, ValueError, BadZipFile, openpyxl.utils.exceptions.InvalidFileException) as exc:
        raise ValueError(f"无法读取 Excel 文件：{exc}") from exc
    try:
        ws = wb.active
        if max(columns) > ws.max_column:
            raise ValueError("配置的列超出 Excel 实际列数")
        students = {}
        for number, row in enumerate(ws.iter_rows(min_row=2, values_only=True), 2):
            value, name, email = [row[c - 1] for c in columns]
            if value is None:
                continue
            if isinstance(value, float) and value.is_integer():
                value = int(value)
            student_id = str(value).strip()
            if not student_id:
                continue
            if student_id in students:
                raise ValueError(f"Excel 第 {number} 行学号重复: {student_id}")
            students[student_id] = {"name": str(name or "").strip(), "email": str(email or "").strip()}
        return students
    finally:
        wb.close()


def read_student_info(excel_path, student_id, student_id_col=1, name_col=2, email_col=3):
    """按学号读取学生信息；批量处理请使用 load_students。"""
    return load_students(excel_path, student_id_col, name_col, email_col).get(str(student_id))


def load_email_template(template_path):
    """读取 UTF-8 模板，第一行是主题，其余为正文。"""
    content = Path(clean_path(template_path)).read_text(encoding="utf-8-sig").strip()
    lines = content.splitlines()
    if not lines:
        raise ValueError("邮件模板为空")
    subject = re.sub(r"^Subject:\s*", "", lines[0], flags=re.I).strip()
    body = "\n".join(lines[1:]).strip()
    if not subject or not body:
        raise ValueError("模板必须包含主题和正文，并以换行分隔")
    return {"subject": subject, "body": body}


def replace_placeholder(text, student_name, student_id=""):
    """替换 {姓名}、{学号}，兼容旧版裸写的“姓名”占位符。"""
    return text.replace("{学号}", str(student_id)).replace("{姓名}", "姓名").replace("姓名", str(student_name))


def file_digest(path):
    """分块计算附件指纹。"""
    digest = hashlib.sha256()
    with open(path, "rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def create_empty_status():
    """创建新版发送日志。"""
    return {"version": 2, "records": {}}


def load_status(status_file=STATUS_FILE):
    """损坏或旧格式日志明确报错，不静默丢弃发送历史。"""
    path = Path(status_file)
    if not path.exists():
        return create_empty_status()
    with path.open(encoding="utf-8") as stream:
        status = json.load(stream)
    if not isinstance(status, dict) or status.get("version") != 2 or not isinstance(status.get("records"), dict):
        raise ValueError("发送日志为旧版或格式损坏。请先备份并核对已发送邮件，再使用新的 --state 文件。")
    for record in status["records"].values():
        if not isinstance(record, dict) or record.get("status") not in {"success", "failed", "sending", "uncertain"}:
            raise ValueError("发送日志包含无效记录，请先备份核对")
    return status


def save_status(status_file, status):
    """持久化状态，写入失败必须停止发送。"""
    atomic_json(status_file, status)


@contextmanager
def status_lock(status_file):
    """持有操作系统锁，避免多个进程同时使用同一日志发送。"""
    path = Path(str(status_file) + ".lock")
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("a+b") as stream:
        stream.seek(0, 2)
        if stream.tell() == 0:
            stream.write(b"0")
            stream.flush()
        stream.seek(0)
        try:
            if os.name == "nt":
                import msvcrt
                msvcrt.locking(stream.fileno(), msvcrt.LK_NBLCK, 1)
            else:
                import fcntl
                fcntl.flock(stream, fcntl.LOCK_EX | fcntl.LOCK_NB)
        except OSError as exc:
            raise RuntimeError("另一个发送任务正在使用这份日志，请等待其结束") from exc
        try:
            yield
        finally:
            stream.seek(0)
            if os.name == "nt":
                msvcrt.locking(stream.fileno(), msvcrt.LK_UNLCK, 1)
            else:
                fcntl.flock(stream, fcntl.LOCK_UN)


def prepare_batch(config, folder_path, status_file=STATUS_FILE):
    """离线预检附件、学生和模板，生成可预览的发送计划。"""
    validate_config(config)
    students = load_students(config["student_info_excel"], *[config.get(k, v) for k, v in
        (("excel_student_id_col", 1), ("excel_name_col", 2), ("excel_email_col", 3))])
    template = load_email_template(config["email_template"])
    status = load_status(status_file)
    jobs = []
    for item in get_files_with_numbers(folder_path):
        job = dict(item, name="", email="", subject="", body="", status="ready", error="", key="")
        try:
            student_id = item["student_id"]
            if not student_id:
                raise ValueError("文件名无法提取学号")
            student = students.get(student_id)
            if not student:
                raise ValueError("Excel 中找不到此学号")
            job.update(student)
            if not job["name"]:
                raise ValueError("姓名为空")
            if not valid_email(job["email"]):
                raise ValueError("收件人邮箱格式无效")
            with open(job["filepath"], "rb") as stream:
                if stream.read(5) != b"%PDF-":
                    raise ValueError("文件缺少 PDF 标识，请检查附件")
            job["digest"] = file_digest(job["filepath"])
            for key in ("subject", "body"):
                job[key] = replace_placeholder(template[key], job["name"], student_id)
            identity = [str(Path(job["filepath"]).resolve()), job["digest"],
                        config.get("sender_email", "").strip().lower(), job["email"], job["subject"], job["body"]]
            job["key"] = hashlib.sha256(json.dumps(identity, ensure_ascii=False).encode()).hexdigest()
            old = status["records"].get(job["key"], {})
            job["status"] = old.get("status", "ready")
            job["error"] = old.get("error", "")
            if job["status"] == "sending":
                job.update(status="uncertain", error="上次发送中断，请核实已发送邮件；此项不会自动重发")
        except (ValueError, OSError) as exc:
            job.update(status="blocked", error=str(exc))
        jobs.append(job)
    return jobs


class DeliveryUncertain(Exception):
    """SMTP 提交中断，服务端可能已经收到邮件。"""


def send_email(smtp_config, recipient, subject, body, attachment_path=None, expected_digest=None):
    """通过校验证书的 SSL/STARTTLS 发送；附件缺失时拒绝发送。"""
    validate_config(smtp_config, sending=True)
    if not valid_email(recipient):
        raise ValueError("收件人地址无效")
    if not attachment_path or not Path(attachment_path).is_file():
        raise FileNotFoundError("PDF 附件不存在，已取消发送")
    message = EmailMessage()
    message["From"] = smtp_config["sender_email"]
    message["To"] = recipient
    message["Subject"] = subject
    message.set_content(body)
    attachment = Path(attachment_path).read_bytes()
    if not attachment.startswith(b"%PDF-"):
        raise ValueError("文件缺少 PDF 标识")
    if expected_digest and hashlib.sha256(attachment).hexdigest() != expected_digest:
        raise ValueError("附件已在预检后更改，请重新预检")
    message.add_attachment(attachment, maintype="application", subtype="pdf",
                           filename=Path(attachment_path).name)
    port = int(smtp_config["smtp_port"])
    implicit_ssl = smtp_config.get("use_ssl") or port == 465
    context = ssl.create_default_context()
    server = None
    try:
        if implicit_ssl:
            server = smtplib.SMTP_SSL(smtp_config["smtp_server"], port, timeout=30, context=context)
        else:
            server = smtplib.SMTP(smtp_config["smtp_server"], port, timeout=30)
            server.ehlo()
            server.starttls(context=context)
            server.ehlo()
        server.login(smtp_config["sender_email"], smtp_config["sender_password"])
        try:
            refused = server.send_message(message)
            if refused:
                raise smtplib.SMTPRecipientsRefused(refused)
        except (smtplib.SMTPServerDisconnected, OSError) as exc:
            raise DeliveryUncertain("提交邮件时连接中断，可能已送达，请核实后再处理") from exc
        return True
    finally:
        if server:
            try:
                server.quit()
            except (smtplib.SMTPException, OSError):
                try:
                    server.close()
                except OSError:
                    pass  # 清理连接失败不能改变已经明确收到的提交结果。


def send_email_with_retry(smtp_config, recipient, subject, body, attachment_path,
                          max_retries=3, retry_interval=10, stop=None, expected_digest=None):
    """仅重试连接故障及 SMTP 临时错误，结果不明确时停止自动重试。"""
    stop = stop or threading.Event()
    for attempt in range(max_retries):
        if stop.is_set():
            return False, "任务已停止，可稍后重试"
        try:
            send_email(smtp_config, recipient, subject, body, attachment_path, expected_digest)
            return True, None
        except DeliveryUncertain:
            raise
        except (OSError, smtplib.SMTPException) as exc:
            network_error = isinstance(exc, OSError) and not isinstance(
                exc, (smtplib.SMTPException, FileNotFoundError, PermissionError, ssl.SSLError))
            transient = network_error or isinstance(exc, smtplib.SMTPServerDisconnected) or 400 <= getattr(exc, "smtp_code", 0) < 500
            if not transient or attempt == max_retries - 1:
                return False, str(exc)
            if stop.wait(retry_interval):
                return False, "任务已停止，可稍后重试"
        except (ValueError, KeyError) as exc:
            return False, str(exc)
    return False, "尝试次数无效"


def run_batch(config, jobs, status_file=STATUS_FILE, stop=None, on_update=None, retry_only=False):
    """按预检快照发送，每封前后写日志，支持在间隔和重试期间停止。"""
    validate_config(config, sending=True)
    stop = stop or threading.Event()
    eligible = {"failed"} if retry_only else {"ready", "failed"}
    with status_lock(status_file):
        status = load_status(status_file)
        queue = [j for j in jobs if j["status"] in eligible]
        for index, job in enumerate(queue):
            if stop.is_set():
                break
            existing = status["records"].get(job["key"], {})
            if existing.get("status") in {"success", "sending", "uncertain"}:
                job.update(status=existing["status"], error=existing.get("error", ""))
                if on_update:
                    on_update(job, index + 1, len(queue))
                continue
            try:
                if file_digest(job["filepath"]) != job["digest"]:
                    raise ValueError("附件已在预检后更改，请重新预检")
            except (ValueError, OSError) as exc:
                job.update(status="failed", error=str(exc))
            else:
                record = {k: job[k] for k in ("filename", "student_id", "name", "email")}
                record.update(status="sending", error="", time=datetime.now().isoformat(timespec="seconds"))
                status["records"][job["key"]] = record
                save_status(status_file, status)
                try:
                    success, error = send_email_with_retry(config, job["email"], job["subject"], job["body"],
                        job["filepath"], int(config.get("max_retries", 3)), float(config.get("retry_interval", 10)), stop, job["digest"])
                    job.update(status="success" if success else "failed", error=error or "")
                except DeliveryUncertain as exc:
                    job.update(status="uncertain", error=str(exc))
            record = {k: job[k] for k in ("filename", "student_id", "name", "email", "status", "error")}
            record["time"] = datetime.now().isoformat(timespec="seconds")
            status["records"][job["key"]] = record
            save_status(status_file, status)
            if on_update:
                on_update(job, index + 1, len(queue))
            if index < len(queue) - 1 and stop.wait(float(config.get("send_interval", 3))):
                break
    return jobs


def export_results(jobs, output_path):
    """导出 XLSX 或带 BOM 的 CSV，防止表格软件将内容当作公式。"""
    keys = ("filename", "student_id", "name", "email", "status", "error")
    rows = [["文件名", "学号", "姓名", "邮箱", "状态", "说明"]]
    for job in jobs:
        values = [LABELS.get(job.get(k), job.get(k, "")) if k == "status" else job.get(k, "") for k in keys]
        rows.append(["'" + str(v) if str(v).startswith(("=", "+", "-", "@", "\t", "\r")) else v for v in values])
    if Path(output_path).suffix.lower() == ".xlsx":
        book = openpyxl.Workbook()
        sheet = book.active
        sheet.title = "检查与发送结果"
        for row in rows:
            sheet.append(row)
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for index, width in enumerate((36, 20, 18, 36, 14, 65), 1):
            sheet.column_dimensions[openpyxl.utils.get_column_letter(index)].width = width
        for cell in sheet[1]:
            cell.font = openpyxl.styles.Font(color="FFFFFF", bold=True)
            cell.fill = openpyxl.styles.PatternFill("solid", fgColor="0067C0")
        try:
            book.save(output_path)
        finally:
            book.close()
        return
    with open(output_path, "w", newline="", encoding="utf-8-sig") as stream:
        writer = csv.writer(stream)
        writer.writerows(rows)


def main(argv=None):
    """命令行支持离线预检、失败重试、指定配置和结果导出。"""
    parser = argparse.ArgumentParser(description="PDF 批量邮件 · 先预检，再发送")
    parser.add_argument("-p", "--path", help="PDF 附件目录")
    parser.add_argument("--config", default="config.json", help="配置文件")
    parser.add_argument("--state", default=STATUS_FILE, help="发送日志路径")
    parser.add_argument("--dry-run", action="store_true", help="只预检，不连接 SMTP")
    parser.add_argument("--retry-failed", action="store_true", help="只重试失败记录")
    parser.add_argument("--export", help="导出 XLSX 或 CSV 结果")
    parser.add_argument("-s", "--status", action="store_true", help="查看日志")
    parser.add_argument("--smtp")
    parser.add_argument("--port", type=int)
    parser.add_argument("--email")
    parser.add_argument("-r", "--restart", action="store_true", help="提示改用新日志，避免重复发送")
    args = parser.parse_args(argv)
    try:
        if args.restart:
            raise ValueError("重新发送请先核对历史，再用 --state 指定新日志；不会自动删除历史")
        if args.status:
            status = load_status(args.state)
            for key, label in LABELS.items():
                print(f"{label}: {sum(r['status'] == key for r in status['records'].values())}")
            return 0
        config = load_config(args.config)
        for arg, key in ((args.smtp, "smtp_server"), (args.port, "smtp_port"), (args.email, "sender_email")):
            if arg is not None:
                config[key] = arg
        folder = args.path or config.get("attachment_folder") or clean_path(input("PDF 附件目录: "))
        if not args.dry_run and not config.get("sender_email"):
            config["sender_email"] = input("发件人邮箱: ").strip()
        jobs = prepare_batch(config, folder, args.state)
        for job in jobs:
            print(f"{LABELS[job['status']]} | {job['filename']} | {job['email']} | {job['error']}")
        if not jobs:
            print("未找到 PDF 文件")
            return 1
        if not args.dry_run:
            eligible = {"failed"} if args.retry_failed else {"ready", "failed"}
            count = sum(j["status"] in eligible for j in jobs)
            if count:
                if not config.get("sender_password"):
                    config["sender_password"] = getpass.getpass("邮箱授权码（不保存）: ")
                if input(f"将发送 {count} 封邮件，输入 SEND 确认: ").strip() == "SEND":
                    run_batch(config, jobs, args.state, retry_only=args.retry_failed,
                              on_update=lambda j, i, n: print(f"[{i}/{n}] {LABELS[j['status']]} {j['filename']} {j['error']}"))
        if args.export:
            export_results(jobs, args.export)
        return 1 if any(j["status"] in {"blocked", "failed", "uncertain", "sending"} for j in jobs) else 0
    except KeyboardInterrupt:
        print("\n已停止；下次运行会依据日志续发。发送中断项需核实。")
        return 130
    except (ValueError, OSError, RuntimeError, smtplib.SMTPException) as exc:
        print(f"错误: {exc}")
        return 1


if __name__ == "__main__":
    raise SystemExit(main())
