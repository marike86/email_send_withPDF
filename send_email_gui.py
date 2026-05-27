import os
import re
import json
import time
import threading
import logging
import smtplib
import webbrowser
from datetime import datetime
from email.mime.multipart import MIMEMultipart
from email.mime.text import MIMEText
from email.mime.application import MIMEApplication
from email.header import Header
from tkinter import filedialog, messagebox, END, WORD, BOTH, LEFT, RIGHT, X, Y, TOP, BOTTOM, W, E, N, S, VERTICAL, HORIZONTAL

import tkinter as tk
from tkinter import ttk
from tkinter.scrolledtext import ScrolledText
import openpyxl

try:
    import customtkinter as ctk
    USE_CUSTOMTK = True
except ImportError:
    USE_CUSTOMTK = False

if USE_CUSTOMTK:
    ctk.set_appearance_mode("light")
    ctk.set_default_color_theme("blue")

PRIMARY = "#0078D4"
PRIMARY_HOVER = "#106EBE"
PRIMARY_PRESSED = "#005A9E"
SUCCESS = "#107C10"
SUCCESS_HOVER = "#0B6A0B"
WARNING = "#FF8C00"
ERROR = "#D13438"
BG = "#F3F3F3"
CARD = "#FFFFFF"
CARD_BORDER = "#E0E0E0"
TEXT_PRIMARY = "#1A1A1A"
TEXT_SECONDARY = "#616161"
TEXT_DISABLED = "#A0A0A0"
INPUT_BG = "#FFFFFF"
INPUT_BORDER = "#8A8A8A"
INPUT_FOCUS = PRIMARY
BORDER_RADIUS = 8
FONT_FAMILY = "Segoe UI"
FONT_SIZE_TITLE = 24
FONT_SIZE_SUBTITLE = 16
FONT_SIZE_BODY = 13
FONT_SIZE_SMALL = 11
FONT_SIZE_BUTTON = 13
FONT_SIZE_STEP_NUM = 16
FONT_SIZE_STEP_LABEL = 10

logging.basicConfig(level=logging.INFO, format="%(asctime)s [%(levelname)s] %(message)s")
logger = logging.getLogger(__name__)


def extract_key_from_filename(filename):
    name = os.path.splitext(filename)[0]
    patterns = [
        r'签署后[：:](\d+)[-_]',
        r'(\d{6,})[-_]',
        r'(\d{6,})',
    ]
    for pat in patterns:
        m = re.search(pat, name)
        if m:
            return m.group(1)
    return None


def suggest_field_type(header):
    h = header.strip().lower()
    email_kw = ["邮箱", "email", "e-mail", "邮件", "mail", "电子邮件"]
    name_kw = ["姓名", "名字", "name", "学生名", "收件人"]
    id_kw = ["学号", "id", "工号", "编号", "studentid", "student_id", "no", "序号"]
    phone_kw = ["电话", "手机", "phone", "tel", "mobile"]
    for kw in email_kw:
        if kw in h:
            return "email"
    for kw in name_kw:
        if kw in h:
            return "name"
    for kw in id_kw:
        if kw in h:
            return "id"
    for kw in phone_kw:
        if kw in h:
            return "phone"
    return "other"


class ColumnMapping:
    def __init__(self):
        self.headers = []
        self.data = []
        self.mapping = {}
        self._col_cache = {}

    def load_from_excel(self, filepath, sheet_name=None):
        wb = openpyxl.load_workbook(filepath, read_only=True, data_only=True)
        ws = wb[sheet_name] if sheet_name and sheet_name in wb.sheetnames else wb.active
        rows = list(ws.iter_rows(values_only=True))
        wb.close()
        if not rows:
            return
        self.headers = [str(h).strip() if h else f"列{i+1}" for i, h in enumerate(rows[0])]
        self.data = rows[1:]
        self._col_cache = {}

    def set_mapping(self, field, col_index):
        self.mapping[field] = col_index
        self._col_cache.pop(field, None)

    def get_key_column(self):
        for field in ("id", "email", "name"):
            if field in self.mapping:
                return self.mapping[field]
        return 0

    def get_key_field_name(self):
        for field in ("id", "email", "name"):
            if field in self.mapping:
                return field
        return "id"

    def get_column_index(self, field):
        return self.mapping.get(field, -1)

    def get_data_for_field(self, field):
        if field in self._col_cache:
            return self._col_cache[field]
        idx = self.get_column_index(field)
        if idx < 0:
            return []
        values = []
        for row in self.data:
            if idx < len(row) and row[idx] is not None:
                values.append(str(row[idx]).strip())
            else:
                values.append("")
        self._col_cache[field] = values
        return values


class RecipientMatcher:
    def __init__(self, folder_path, column_mapping, key_field="id"):
        self.folder_path = folder_path
        self.column_mapping = column_mapping
        self.key_field = key_field
        self._key_index = {}
        self._load_recipients()

    def _load_recipients(self):
        key_idx = self.column_mapping.get_column_index(self.key_field)
        if key_idx < 0:
            return
        for i, row in enumerate(self.column_mapping.data):
            if key_idx < len(row) and row[key_idx] is not None:
                raw = str(row[key_idx]).strip()
                if "." in raw and raw.replace(".", "").isdigit():
                    raw = raw.split(".")[0]
                self._key_index[raw] = i

    def find_by_key(self, key):
        return self._key_index.get(key)

    def match_files(self):
        results = []
        if not os.path.isdir(self.folder_path):
            return results
        for fn in sorted(os.listdir(self.folder_path)):
            fp = os.path.join(self.folder_path, fn)
            if not os.path.isfile(fp):
                continue
            key = extract_key_from_filename(fn)
            if key is None:
                continue
            row_idx = self.find_by_key(key)
            info = {"filename": fn, "filepath": fp, "key": key, "row_idx": row_idx, "matched": row_idx is not None}
            results.append(info)
        return results


class EmailSenderGUI:
    def __init__(self):
        self.root = ctk.CTk() if USE_CUSTOMTK else tk.Tk()
        self.root.title("批量邮件发送工具")
        self.root.geometry("900x700")
        self.root.minsize(800, 600)
        self.root.configure(fg_color=BG if USE_CUSTOMTK else BG)
        self._configure_dpi()
        self.current_step = 0
        self.step_frames = []
        self.step_buttons = []
        self.folder_path = tk.StringVar(value="")
        self.excel_path = tk.StringVar(value="")
        self.smtp_server = tk.StringVar(value="smtp.gmail.com")
        self.smtp_port = tk.IntVar(value=587)
        self.use_tls = tk.BooleanVar(value=True)
        self.sender_email = tk.StringVar(value="")
        self.sender_password = tk.StringVar(value="")
        self.send_interval = tk.DoubleVar(value=3.0)
        self.max_retries = tk.IntVar(value=3)
        self.retry_interval = tk.IntVar(value=10)
        self.subject_var = tk.StringVar(value="")
        self.body_widget = None
        self.log_widget = None
        self.progress_bar = None
        self.progress_label = None
        self.column_mapping = ColumnMapping()
        self.mapping_widgets = {}
        self.stop_event = threading.Event()
        self.pause_event = threading.Event()
        self.pause_event.set()
        self.is_sending = False
        self.send_thread = None
        self.total_count = 0
        self.success_count = 0
        self.fail_count = 0
        self.failed_records = []
        self.custom_vars = []
        self.custom_var_container = None
        self.custom_var_subject_container = None
        self.custom_var_body_container = None
        self._build_ui()
        self.root.protocol("WM_DELETE_WINDOW", self.on_closing)

    def _configure_dpi(self):
        try:
            from ctypes import windll
            windll.shcore.SetProcessDpiAwareness(1)
        except Exception:
            pass

    # ------------------------------------------------------------------ UI
    def _build_ui(self):
        self.nav_frame = self._create_nav_bar()
        self.button_frame = self._create_bottom_buttons()
        self.content_frame = self._create_content_area()
        self._create_all_steps()
        self.show_step(0)

    def _create_nav_bar(self):
        if USE_CUSTOMTK:
            frame = ctk.CTkFrame(self.root, fg_color=CARD, corner_radius=0, height=72)
        else:
            frame = tk.Frame(self.root, bg=CARD, height=72)
            frame.pack_propagate(False)
        frame.pack(fill=X, side=TOP)

        if USE_CUSTOMTK:
            title = ctk.CTkLabel(frame, text="批量邮件发送", font=(FONT_FAMILY, FONT_SIZE_TITLE, "bold"), text_color=PRIMARY)
        else:
            title = tk.Label(frame, text="批量邮件发送", font=(FONT_FAMILY, FONT_SIZE_TITLE, "bold"), fg=PRIMARY, bg=CARD)
        title.pack(side=LEFT, padx=(24, 0), pady=12)

        steps_container = ctk.CTkFrame(frame, fg_color="transparent") if USE_CUSTOMTK else tk.Frame(frame, bg=CARD)
        steps_container.pack(side=RIGHT, padx=(0, 24), pady=12)

        step_names = ["选择文件", "编辑邮件", "SMTP 设置", "发送邮件"]
        self.step_buttons = []
        for i, name in enumerate(step_names):
            btn = self._create_step_button(steps_container, i, name)
            self.step_buttons.append(btn)
            if i < len(step_names) - 1:
                if USE_CUSTOMTK:
                    sep = ctk.CTkLabel(steps_container, text="—", text_color=TEXT_DISABLED, font=(FONT_FAMILY, 10))
                else:
                    sep = tk.Label(steps_container, text="—", fg=TEXT_DISABLED, bg=CARD, font=(FONT_FAMILY, 10))
                sep.pack(side=LEFT, padx=4)
        return frame

    def _create_step_button(self, parent, index, label):
        if USE_CUSTOMTK:
            btn = ctk.CTkButton(
                parent, text=f"{index+1}", width=36, height=36,
                corner_radius=18, fg_color=TEXT_DISABLED, hover_color=PRIMARY_HOVER,
                font=(FONT_FAMILY, FONT_SIZE_STEP_NUM, "bold"),
                command=lambda: self.go_to_step(index)
            )
        else:
            btn = tk.Button(
                parent, text=f"{index+1}", width=3, height=1,
                bg=TEXT_DISABLED, fg="white", font=(FONT_FAMILY, FONT_SIZE_STEP_NUM, "bold"),
                relief="flat", activebackground=PRIMARY_HOVER, activeforeground="white",
                command=lambda: self.go_to_step(index)
            )
        btn.pack(side=LEFT, padx=2)

        if USE_CUSTOMTK:
            lbl = ctk.CTkLabel(parent, text=label, font=(FONT_FAMILY, FONT_SIZE_STEP_LABEL), text_color=TEXT_SECONDARY)
        else:
            lbl = tk.Label(parent, text=label, font=(FONT_FAMILY, FONT_SIZE_STEP_LABEL), fg=TEXT_SECONDARY, bg=CARD)
        lbl.pack(side=LEFT, padx=(2, 8))
        btn._label_ref = lbl
        return btn

    def _update_step_buttons(self):
        for i, btn in enumerate(self.step_buttons):
            if i == self.current_step:
                if USE_CUSTOMTK:
                    btn.configure(fg_color=PRIMARY, text_color="white")
                    btn._label_ref.configure(text_color=PRIMARY)
                else:
                    btn.configure(bg=PRIMARY, fg="white")
                    btn._label_ref.configure(fg=PRIMARY)
            elif i < self.current_step:
                if USE_CUSTOMTK:
                    btn.configure(fg_color=SUCCESS, text_color="white")
                    btn._label_ref.configure(text_color=SUCCESS)
                else:
                    btn.configure(bg=SUCCESS, fg="white")
                    btn._label_ref.configure(fg=SUCCESS)
            else:
                if USE_CUSTOMTK:
                    btn.configure(fg_color=TEXT_DISABLED, text_color="white")
                    btn._label_ref.configure(text_color=TEXT_SECONDARY)
                else:
                    btn.configure(bg=TEXT_DISABLED, fg="white")
                    btn._label_ref.configure(fg=TEXT_SECONDARY)

    def _create_content_area(self):
        if USE_CUSTOMTK:
            frame = ctk.CTkFrame(self.root, fg_color=BG, corner_radius=0)
        else:
            frame = tk.Frame(self.root, bg=BG)
        frame.pack(fill=BOTH, expand=True, padx=0, pady=0)
        return frame

    def _create_bottom_buttons(self):
        if USE_CUSTOMTK:
            frame = ctk.CTkFrame(self.root, fg_color=CARD, corner_radius=0, height=64)
        else:
            frame = tk.Frame(self.root, bg=CARD, height=64)
            frame.pack_propagate(False)
        frame.pack(fill=X, side=BOTTOM)

        inner = ctk.CTkFrame(frame, fg_color="transparent") if USE_CUSTOMTK else tk.Frame(frame, bg=CARD)
        inner.pack(fill=X, padx=24, pady=12)

        if USE_CUSTOMTK:
            self.prev_btn = ctk.CTkButton(inner, text="上一步", width=100, height=36, corner_radius=BORDER_RADIUS,
                                          fg_color="transparent", border_width=1, border_color=INPUT_BORDER,
                                          text_color=TEXT_PRIMARY, hover_color="#E8E8E8", font=(FONT_FAMILY, FONT_SIZE_BUTTON),
                                          command=self.prev_step)
        else:
            self.prev_btn = tk.Button(inner, text="上一步", width=10, height=1, relief="flat",
                                      bg=CARD, fg=TEXT_PRIMARY, font=(FONT_FAMILY, FONT_SIZE_BUTTON),
                                      activebackground="#E8E8E8", command=self.prev_step)
        self.prev_btn.pack(side=LEFT)

        if USE_CUSTOMTK:
            self.next_btn = ctk.CTkButton(inner, text="下一步", width=100, height=36, corner_radius=BORDER_RADIUS,
                                          fg_color=PRIMARY, hover_color=PRIMARY_HOVER, text_color="white",
                                          font=(FONT_FAMILY, FONT_SIZE_BUTTON), command=self.next_step)
        else:
            self.next_btn = tk.Button(inner, text="下一步", width=10, height=1, relief="flat",
                                      bg=PRIMARY, fg="white", font=(FONT_FAMILY, FONT_SIZE_BUTTON),
                                      activebackground=PRIMARY_HOVER, activeforeground="white", command=self.next_step)
        self.next_btn.pack(side=RIGHT)

        if USE_CUSTOMTK:
            self.send_btn = ctk.CTkButton(inner, text="开始发送", width=120, height=36, corner_radius=BORDER_RADIUS,
                                          fg_color=SUCCESS, hover_color=SUCCESS_HOVER, text_color="white",
                                          font=(FONT_FAMILY, FONT_SIZE_BUTTON, "bold"), command=self.start_send)
        else:
            self.send_btn = tk.Button(inner, text="开始发送", width=12, height=1, relief="flat",
                                      bg=SUCCESS, fg="white", font=(FONT_FAMILY, FONT_SIZE_BUTTON, "bold"),
                                      activebackground=SUCCESS_HOVER, activeforeground="white", command=self.start_send)

        return frame

    # ----------------------------------------------------------- steps
    def _create_all_steps(self):
        self.step_frames = [self.create_step1(), self.create_step2(), self.create_step3(), self.create_step4()]

    def create_step1(self):
        if USE_CUSTOMTK:
            frame = ctk.CTkFrame(self.content_frame, fg_color=BG)
        else:
            frame = tk.Frame(self.content_frame, bg=BG)
        self._card(frame, "附件文件夹", self._build_step1_folder)
        self._card(frame, "Excel 信息表", self._build_step1_excel)
        self._card(frame, "列映射配置", self._build_step1_mapping)
        self._card(frame, "匹配预览", self._build_step1_preview)
        return frame

    def _build_step1_folder(self, parent):
        row = ctk.CTkFrame(parent, fg_color="transparent") if USE_CUSTOMTK else tk.Frame(parent, bg=CARD)
        row.pack(fill=X, pady=(0, 4))
        if USE_CUSTOMTK:
            entry = ctk.CTkEntry(row, textvariable=self.folder_path, placeholder_text="请选择包含附件的文件夹路径",
                                 height=36, border_width=1, border_color=INPUT_BORDER, fg_color=INPUT_BG,
                                 text_color=TEXT_PRIMARY, font=(FONT_FAMILY, FONT_SIZE_BODY))
            entry.pack(side=LEFT, fill=X, expand=True, padx=(0, 8))
            ctk.CTkButton(row, text="浏览…", width=80, height=36, corner_radius=BORDER_RADIUS,
                          fg_color=PRIMARY, hover_color=PRIMARY_HOVER, text_color="white",
                          font=(FONT_FAMILY, FONT_SIZE_BODY), command=self.browse_folder).pack(side=RIGHT)
        else:
            entry = tk.Entry(row, textvariable=self.folder_path, font=(FONT_FAMILY, FONT_SIZE_BODY), relief="flat",
                             bg=INPUT_BG, fg=TEXT_PRIMARY, highlightthickness=1, highlightcolor=INPUT_FOCUS,
                             highlightbackground=INPUT_BORDER)
            entry.pack(side=LEFT, fill=X, expand=True, padx=(0, 8))
            tk.Button(row, text="浏览…", bg=PRIMARY, fg="white", relief="flat", font=(FONT_FAMILY, FONT_SIZE_BODY),
                      activebackground=PRIMARY_HOVER, command=self.browse_folder).pack(side=RIGHT)

    def _build_step1_excel(self, parent):
        row = ctk.CTkFrame(parent, fg_color="transparent") if USE_CUSTOMTK else tk.Frame(parent, bg=CARD)
        row.pack(fill=X, pady=(0, 4))
        if USE_CUSTOMTK:
            entry = ctk.CTkEntry(row, textvariable=self.excel_path, placeholder_text="请选择 Excel 文件路径",
                                 height=36, border_width=1, border_color=INPUT_BORDER, fg_color=INPUT_BG,
                                 text_color=TEXT_PRIMARY, font=(FONT_FAMILY, FONT_SIZE_BODY))
            entry.pack(side=LEFT, fill=X, expand=True, padx=(0, 8))
            ctk.CTkButton(row, text="浏览…", width=80, height=36, corner_radius=BORDER_RADIUS,
                          fg_color=PRIMARY, hover_color=PRIMARY_HOVER, text_color="white",
                          font=(FONT_FAMILY, FONT_SIZE_BODY), command=self.browse_excel).pack(side=RIGHT)
            ctk.CTkButton(row, text="加载列", width=80, height=36, corner_radius=BORDER_RADIUS,
                          fg_color=WARNING, hover_color="#E07B00", text_color="white",
                          font=(FONT_FAMILY, FONT_SIZE_BODY), command=self.load_columns).pack(side=RIGHT, padx=(0, 8))
        else:
            entry = tk.Entry(row, textvariable=self.excel_path, font=(FONT_FAMILY, FONT_SIZE_BODY), relief="flat",
                             bg=INPUT_BG, fg=TEXT_PRIMARY, highlightthickness=1, highlightcolor=INPUT_FOCUS,
                             highlightbackground=INPUT_BORDER)
            entry.pack(side=LEFT, fill=X, expand=True, padx=(0, 8))
            tk.Button(row, text="加载列", bg=WARNING, fg="white", relief="flat", font=(FONT_FAMILY, FONT_SIZE_BODY),
                      activebackground="#E07B00", command=self.load_columns).pack(side=RIGHT, padx=(0, 8))
            tk.Button(row, text="浏览…", bg=PRIMARY, fg="white", relief="flat", font=(FONT_FAMILY, FONT_SIZE_BODY),
                      activebackground=PRIMARY_HOVER, command=self.browse_excel).pack(side=RIGHT)

    def _build_step1_mapping(self, parent):
        self.mapping_container = ctk.CTkFrame(parent, fg_color="transparent") if USE_CUSTOMTK else tk.Frame(parent, bg=CARD)
        self.mapping_container.pack(fill=X)
        if USE_CUSTOMTK:
            ctk.CTkLabel(self.mapping_container, text="请先加载 Excel 文件以配置列映射",
                         font=(FONT_FAMILY, FONT_SIZE_BODY), text_color=TEXT_SECONDARY).pack(pady=8)
        else:
            tk.Label(self.mapping_container, text="请先加载 Excel 文件以配置列映射",
                     font=(FONT_FAMILY, FONT_SIZE_BODY), fg=TEXT_SECONDARY, bg=CARD).pack(pady=8)

    def _build_step1_preview(self, parent):
        self.preview_hint = None
        if USE_CUSTOMTK:
            self.preview_hint = ctk.CTkLabel(parent, text="请先选择文件夹和 Excel 文件，然后加载列映射",
                                             font=(FONT_FAMILY, FONT_SIZE_SMALL), text_color=TEXT_SECONDARY)
            self.preview_hint.pack(fill=X, pady=(0, 4))
            btn_row = ctk.CTkFrame(parent, fg_color="transparent")
            btn_row.pack(fill=X, pady=(0, 4))
            ctk.CTkButton(btn_row, text="加载预览", width=100, height=32, corner_radius=BORDER_RADIUS,
                          fg_color=PRIMARY, hover_color=PRIMARY_HOVER, text_color="white",
                          font=(FONT_FAMILY, FONT_SIZE_BODY), command=self.load_preview).pack(side=LEFT)
            self.preview_container = ctk.CTkScrollableFrame(parent, height=160, fg_color=BG,
                                                             border_width=1, border_color=CARD_BORDER)
            self.preview_container.pack(fill=BOTH, expand=True)
        else:
            self.preview_hint = tk.Label(parent, text="请先选择文件夹和 Excel 文件，然后加载列映射",
                                         font=(FONT_FAMILY, FONT_SIZE_SMALL), fg=TEXT_SECONDARY, bg=CARD)
            self.preview_hint.pack(fill=X, pady=(0, 4))
            btn_row = tk.Frame(parent, bg=CARD)
            btn_row.pack(fill=X, pady=(0, 4))
            tk.Button(btn_row, text="加载预览", bg=PRIMARY, fg="white", relief="flat",
                      font=(FONT_FAMILY, FONT_SIZE_BODY), activebackground=PRIMARY_HOVER,
                      command=self.load_preview).pack(side=LEFT)
            self.preview_container = tk.Frame(parent, bg=CARD)
            self.preview_container.pack(fill=BOTH, expand=True)

    def _build_mapping_widgets(self):
        for w in self.mapping_container.winfo_children():
            w.destroy()
        self.mapping_widgets = {}
        if not self.column_mapping.headers:
            return
        fields = [("id", "关键列（匹配附件）"), ("email", "邮箱列"), ("name", "姓名列")]
        headers = self.column_mapping.headers
        for field, label in fields:
            row = ctk.CTkFrame(self.mapping_container, fg_color="transparent") if USE_CUSTOMTK else tk.Frame(self.mapping_container, bg=CARD)
            row.pack(fill=X, pady=2)
            if USE_CUSTOMTK:
                ctk.CTkLabel(row, text=label, width=160, anchor="w", font=(FONT_FAMILY, FONT_SIZE_BODY),
                             text_color=TEXT_PRIMARY).pack(side=LEFT)
                combo = ctk.CTkComboBox(row, values=headers, width=240, height=32,
                                        border_width=1, border_color=INPUT_BORDER, fg_color=INPUT_BG,
                                        text_color=TEXT_PRIMARY, font=(FONT_FAMILY, FONT_SIZE_BODY),
                                        dropdown_fg_color=CARD, dropdown_hover_color=BG,
                                        command=lambda val, f=field: self._on_mapping_change(f, val))
            else:
                tk.Label(row, text=label, width=20, anchor="w", font=(FONT_FAMILY, FONT_SIZE_BODY),
                         fg=TEXT_PRIMARY, bg=CARD).pack(side=LEFT)
                combo = ttk.Combobox(row, values=headers, width=30, state="readonly",
                                     font=(FONT_FAMILY, FONT_SIZE_BODY))
                combo.bind("<<ComboboxSelected>>", lambda e, f=field, c=combo: self._on_mapping_change(f, c.get()))
            suggested = self._suggest_column(field, headers)
            if suggested:
                combo.set(suggested)
                idx = headers.index(suggested)
                self.column_mapping.set_mapping(field, idx)
            combo.pack(side=LEFT, padx=8)
            self.mapping_widgets[field] = combo

    def _suggest_column(self, field, headers):
        for i, h in enumerate(headers):
            if suggest_field_type(h) == field:
                return h
        return None

    def _on_mapping_change(self, field, value):
        headers = self.column_mapping.headers
        if value in headers:
            self.column_mapping.set_mapping(field, headers.index(value))

    def _card(self, parent, title, build_fn):
        if USE_CUSTOMTK:
            card = ctk.CTkFrame(parent, fg_color=CARD, corner_radius=BORDER_RADIUS, border_width=1,
                                border_color=CARD_BORDER)
            card.pack(fill=X, padx=24, pady=6)
            ctk.CTkLabel(card, text=title, font=(FONT_FAMILY, FONT_SIZE_SUBTITLE, "bold"),
                         text_color=TEXT_PRIMARY, anchor="w").pack(fill=X, padx=16, pady=(12, 4))
            inner = ctk.CTkFrame(card, fg_color="transparent")
            inner.pack(fill=X, padx=16, pady=(0, 12))
        else:
            card = tk.Frame(parent, bg=CARD, highlightthickness=1, highlightbackground=CARD_BORDER)
            card.pack(fill=X, padx=24, pady=6)
            tk.Label(card, text=title, font=(FONT_FAMILY, FONT_SIZE_SUBTITLE, "bold"),
                     fg=TEXT_PRIMARY, bg=CARD, anchor="w").pack(fill=X, padx=16, pady=(12, 4))
            inner = tk.Frame(card, bg=CARD)
            inner.pack(fill=X, padx=16, pady=(0, 12))
        build_fn(inner)

    def create_step2(self):
        if USE_CUSTOMTK:
            frame = ctk.CTkFrame(self.content_frame, fg_color=BG)
        else:
            frame = tk.Frame(self.content_frame, bg=BG)
        self._card(frame, "邮件主题", self._build_step2_subject)
        self._card(frame, "邮件正文", self._build_step2_body)
        self._card(frame, "插入变量", self._build_step2_vars)
        self._card(frame, "自定义变量", self._build_step2_custom_vars)
        self._build_step2_format_toolbar(frame)
        return frame

    def _build_step2_subject(self, parent):
        if USE_CUSTOMTK:
            entry = ctk.CTkEntry(parent, textvariable=self.subject_var, placeholder_text="请输入邮件主题，可用 {姓名} 等变量",
                                 height=36, border_width=1, border_color=INPUT_BORDER, fg_color=INPUT_BG,
                                 text_color=TEXT_PRIMARY, font=(FONT_FAMILY, FONT_SIZE_BODY))
            entry.pack(fill=X, pady=(0, 4))
        else:
            entry = tk.Entry(parent, textvariable=self.subject_var, font=(FONT_FAMILY, FONT_SIZE_BODY), relief="flat",
                             bg=INPUT_BG, fg=TEXT_PRIMARY, highlightthickness=1, highlightcolor=INPUT_FOCUS,
                             highlightbackground=INPUT_BORDER)
            entry.pack(fill=X, pady=(0, 4))
        var_row = ctk.CTkFrame(parent, fg_color="transparent") if USE_CUSTOMTK else tk.Frame(parent, bg=CARD)
        var_row.pack(fill=X)
        self.custom_var_subject_container = var_row
        if USE_CUSTOMTK:
            ctk.CTkLabel(var_row, text="插入变量:", font=(FONT_FAMILY, FONT_SIZE_SMALL), text_color=TEXT_SECONDARY).pack(side=LEFT, padx=(0, 4))
        else:
            tk.Label(var_row, text="插入变量:", font=(FONT_FAMILY, FONT_SIZE_SMALL), fg=TEXT_SECONDARY, bg=CARD).pack(side=LEFT, padx=(0, 4))
        for var in ["{姓名}", "{学号}", "{邮箱}"]:
            if USE_CUSTOMTK:
                ctk.CTkButton(var_row, text=var, width=80, height=26, corner_radius=12, fg_color=BG,
                              hover_color="#D8D8D8", text_color=PRIMARY, font=(FONT_FAMILY, FONT_SIZE_SMALL),
                              border_width=1, border_color=CARD_BORDER,
                              command=lambda v=var: self.insert_var_to_subject(v)).pack(side=LEFT, padx=2)
            else:
                tk.Button(var_row, text=var, bg=BG, fg=PRIMARY, relief="flat", font=(FONT_FAMILY, FONT_SIZE_SMALL),
                          activebackground="#D8D8D8", command=lambda v=var: self.insert_var_to_subject(v)).pack(side=LEFT, padx=2)

    def _build_step2_body(self, parent):
        self.body_widget = ScrolledText(parent, height=10, font=(FONT_FAMILY, FONT_SIZE_BODY), wrap=WORD,
                                        bg=INPUT_BG, fg=TEXT_PRIMARY, insertbackground=TEXT_PRIMARY,
                                        relief="flat", highlightthickness=1,
                                        highlightcolor=INPUT_FOCUS, highlightbackground=INPUT_BORDER,
                                        padx=8, pady=8)
        self.body_widget.pack(fill=BOTH, expand=True, pady=(0, 4))
        self.body_widget.tag_configure("font_style", font=(FONT_FAMILY, FONT_SIZE_BODY))
        for size_name, size_val in [("size_small", 10), ("size_normal", 13), ("size_large", 16), ("size_xlarge", 20), ("size_xxlarge", 24)]:
            self.body_widget.tag_configure(size_name, font=(FONT_FAMILY, size_val))
        self.body_widget.tag_configure("align_right", justify="right")
        self.body_widget.tag_configure("align_left", justify="left")
        self.body_widget.tag_configure("align_center", justify="center")
        self.body_widget.tag_configure("indented", lmargin1=40, lmargin2=40)
        default_body = "姓名同学：\n\n您好！请查收附件。\n\n祝好！"
        self.body_widget.insert("1.0", default_body)

    def _build_step2_vars(self, parent):
        row = ctk.CTkFrame(parent, fg_color="transparent") if USE_CUSTOMTK else tk.Frame(parent, bg=CARD)
        row.pack(fill=X)
        variables = ["{姓名}", "{学号}", "{邮箱}"]
        for var in variables:
            if USE_CUSTOMTK:
                ctk.CTkButton(row, text=var, width=80, height=28, corner_radius=14, fg_color=BG,
                              hover_color="#D8D8D8", text_color=PRIMARY, font=(FONT_FAMILY, FONT_SIZE_SMALL),
                              border_width=1, border_color=CARD_BORDER,
                              command=lambda v=var: self.insert_var(v)).pack(side=LEFT, padx=4)
            else:
                tk.Button(row, text=var, bg=BG, fg=PRIMARY, relief="flat", font=(FONT_FAMILY, FONT_SIZE_SMALL),
                          activebackground="#D8D8D8", command=lambda v=var: self.insert_var(v)).pack(side=LEFT, padx=4)
        self.custom_var_body_container = row

    def _build_step2_custom_vars(self, parent):
        add_row = ctk.CTkFrame(parent, fg_color="transparent") if USE_CUSTOMTK else tk.Frame(parent, bg=CARD)
        add_row.pack(fill=X, pady=(0, 4))
        if USE_CUSTOMTK:
            ctk.CTkLabel(add_row, text="变量名:", font=(FONT_FAMILY, FONT_SIZE_SMALL),
                         text_color=TEXT_PRIMARY).pack(side=LEFT, padx=(0, 4))
            self.cv_name_entry = ctk.CTkEntry(add_row, width=120, height=30, placeholder_text="如: 班级",
                                              border_width=1, border_color=INPUT_BORDER, fg_color=INPUT_BG,
                                              text_color=TEXT_PRIMARY, font=(FONT_FAMILY, FONT_SIZE_SMALL))
            self.cv_name_entry.pack(side=LEFT, padx=(0, 8))
            ctk.CTkLabel(add_row, text="变量值:", font=(FONT_FAMILY, FONT_SIZE_SMALL),
                         text_color=TEXT_PRIMARY).pack(side=LEFT, padx=(0, 4))
            self.cv_value_entry = ctk.CTkEntry(add_row, width=160, height=30, placeholder_text="如: 计算机2班",
                                               border_width=1, border_color=INPUT_BORDER, fg_color=INPUT_BG,
                                               text_color=TEXT_PRIMARY, font=(FONT_FAMILY, FONT_SIZE_SMALL))
            self.cv_value_entry.pack(side=LEFT, padx=(0, 8))
            ctk.CTkButton(add_row, text="添加", width=60, height=30, corner_radius=BORDER_RADIUS,
                          fg_color=PRIMARY, hover_color=PRIMARY_HOVER, text_color="white",
                          font=(FONT_FAMILY, FONT_SIZE_SMALL), command=self._add_custom_var).pack(side=LEFT, padx=2)
            ctk.CTkButton(add_row, text="预览", width=60, height=30, corner_radius=BORDER_RADIUS,
                          fg_color=WARNING, hover_color="#E07B00", text_color="white",
                          font=(FONT_FAMILY, FONT_SIZE_SMALL), command=self._preview_custom_vars).pack(side=LEFT, padx=2)
        else:
            tk.Label(add_row, text="变量名:", font=(FONT_FAMILY, FONT_SIZE_SMALL),
                     fg=TEXT_PRIMARY, bg=CARD).pack(side=LEFT, padx=(0, 4))
            self.cv_name_entry = tk.Entry(add_row, width=15, font=(FONT_FAMILY, FONT_SIZE_SMALL), relief="flat",
                                          bg=INPUT_BG, fg=TEXT_PRIMARY, highlightthickness=1,
                                          highlightcolor=INPUT_FOCUS, highlightbackground=INPUT_BORDER)
            self.cv_name_entry.pack(side=LEFT, padx=(0, 8))
            tk.Label(add_row, text="变量值:", font=(FONT_FAMILY, FONT_SIZE_SMALL),
                     fg=TEXT_PRIMARY, bg=CARD).pack(side=LEFT, padx=(0, 4))
            self.cv_value_entry = tk.Entry(add_row, width=20, font=(FONT_FAMILY, FONT_SIZE_SMALL), relief="flat",
                                           bg=INPUT_BG, fg=TEXT_PRIMARY, highlightthickness=1,
                                           highlightcolor=INPUT_FOCUS, highlightbackground=INPUT_BORDER)
            self.cv_value_entry.pack(side=LEFT, padx=(0, 8))
            tk.Button(add_row, text="添加", bg=PRIMARY, fg="white", relief="flat", width=6, height=1,
                      font=(FONT_FAMILY, FONT_SIZE_SMALL), activebackground=PRIMARY_HOVER,
                      command=self._add_custom_var).pack(side=LEFT, padx=2)
            tk.Button(add_row, text="预览", bg=WARNING, fg="white", relief="flat", width=6, height=1,
                      font=(FONT_FAMILY, FONT_SIZE_SMALL), activebackground="#E07B00",
                      command=self._preview_custom_vars).pack(side=LEFT, padx=2)
        self.custom_var_container = ctk.CTkFrame(parent, fg_color="transparent") if USE_CUSTOMTK else tk.Frame(parent, bg=CARD)
        self.custom_var_container.pack(fill=X)
        self._refresh_custom_var_list()

    def _add_custom_var(self):
        name = self.cv_name_entry.get().strip()
        value = self.cv_value_entry.get().strip()
        if not name:
            messagebox.showwarning("提示", "请输入变量名")
            return
        if not re.match(r'^[\w\u4e00-\u9fff]+$', name):
            messagebox.showwarning("提示", "变量名只能包含字母、数字、下划线或中文")
            return
        reserved = ["姓名", "学号", "邮箱"]
        if name in reserved:
            messagebox.showwarning("提示", f"变量名 '{name}' 为系统内置变量，请更换名称")
            return
        for cv in self.custom_vars:
            if cv["name"] == name:
                messagebox.showwarning("提示", f"变量 '{name}' 已存在，请勿重复添加")
                return
        if not value:
            value = name
        self.custom_vars.append({"name": name, "value": value})
        self.cv_name_entry.delete(0, END)
        self.cv_value_entry.delete(0, END)
        self._refresh_custom_var_list()
        self._refresh_all_var_buttons()

    def _edit_custom_var(self, index):
        if index < 0 or index >= len(self.custom_vars):
            return
        old = self.custom_vars[index]
        dialog = tk.Toplevel(self.root)
        dialog.title("编辑变量")
        dialog.geometry("350x160")
        dialog.resizable(False, False)
        dialog.transient(self.root)
        dialog.grab_set()
        if USE_CUSTOMTK:
            dialog.configure(fg_color=BG)
        else:
            dialog.configure(bg=BG)
        dlg_frame = ctk.CTkFrame(dialog, fg_color="transparent") if USE_CUSTOMTK else tk.Frame(dialog, bg=BG)
        dlg_frame.pack(fill=BOTH, expand=True, padx=20, pady=16)
        if USE_CUSTOMTK:
            ctk.CTkLabel(dlg_frame, text="变量名:", font=(FONT_FAMILY, FONT_SIZE_BODY),
                         text_color=TEXT_PRIMARY).pack(anchor="w")
            name_entry = ctk.CTkEntry(dlg_frame, width=280, height=32, border_width=1,
                                      border_color=INPUT_BORDER, fg_color=INPUT_BG, text_color=TEXT_PRIMARY,
                                      font=(FONT_FAMILY, FONT_SIZE_BODY))
            name_entry.insert(0, old["name"])
            name_entry.configure(state="disabled")
            name_entry.pack(pady=(4, 8))
            ctk.CTkLabel(dlg_frame, text="变量值:", font=(FONT_FAMILY, FONT_SIZE_BODY),
                         text_color=TEXT_PRIMARY).pack(anchor="w")
            value_entry = ctk.CTkEntry(dlg_frame, width=280, height=32, border_width=1,
                                       border_color=INPUT_BORDER, fg_color=INPUT_BG, text_color=TEXT_PRIMARY,
                                       font=(FONT_FAMILY, FONT_SIZE_BODY))
            value_entry.insert(0, old["value"])
            value_entry.pack(pady=(4, 8))
        else:
            tk.Label(dlg_frame, text="变量名:", font=(FONT_FAMILY, FONT_SIZE_BODY),
                     fg=TEXT_PRIMARY, bg=BG).pack(anchor="w")
            name_entry = tk.Entry(dlg_frame, width=35, font=(FONT_FAMILY, FONT_SIZE_BODY), relief="flat",
                                  bg=INPUT_BG, fg=TEXT_DISABLED, state="disabled",
                                  highlightthickness=1, highlightcolor=INPUT_FOCUS, highlightbackground=INPUT_BORDER)
            name_entry.insert(0, old["name"])
            name_entry.pack(pady=(4, 8))
            tk.Label(dlg_frame, text="变量值:", font=(FONT_FAMILY, FONT_SIZE_BODY),
                     fg=TEXT_PRIMARY, bg=BG).pack(anchor="w")
            value_entry = tk.Entry(dlg_frame, width=35, font=(FONT_FAMILY, FONT_SIZE_BODY), relief="flat",
                                   bg=INPUT_BG, fg=TEXT_PRIMARY, highlightthickness=1,
                                   highlightcolor=INPUT_FOCUS, highlightbackground=INPUT_BORDER)
            value_entry.insert(0, old["value"])
            value_entry.pack(pady=(4, 8))
        def do_save():
            new_value = value_entry.get().strip()
            if not new_value:
                new_value = old["name"]
            self.custom_vars[index]["value"] = new_value
            dialog.destroy()
            self._refresh_custom_var_list()
            self._refresh_all_var_buttons()
        btn_frame = ctk.CTkFrame(dlg_frame, fg_color="transparent") if USE_CUSTOMTK else tk.Frame(dlg_frame, bg=BG)
        btn_frame.pack(fill=X)
        if USE_CUSTOMTK:
            ctk.CTkButton(btn_frame, text="保存", width=80, height=32, corner_radius=BORDER_RADIUS,
                          fg_color=PRIMARY, hover_color=PRIMARY_HOVER, text_color="white",
                          font=(FONT_FAMILY, FONT_SIZE_BODY), command=do_save).pack(side=RIGHT, padx=4)
            ctk.CTkButton(btn_frame, text="取消", width=80, height=32, corner_radius=BORDER_RADIUS,
                          fg_color="transparent", border_width=1, border_color=INPUT_BORDER,
                          text_color=TEXT_PRIMARY, font=(FONT_FAMILY, FONT_SIZE_BODY),
                          command=dialog.destroy).pack(side=RIGHT, padx=4)
        else:
            tk.Button(btn_frame, text="保存", bg=PRIMARY, fg="white", relief="flat", width=8, height=1,
                      font=(FONT_FAMILY, FONT_SIZE_BODY), activebackground=PRIMARY_HOVER,
                      command=do_save).pack(side=RIGHT, padx=4)
            tk.Button(btn_frame, text="取消", bg=CARD, fg=TEXT_PRIMARY, relief="flat", width=8, height=1,
                      font=(FONT_FAMILY, FONT_SIZE_BODY), activebackground="#E8E8E8",
                      command=dialog.destroy).pack(side=RIGHT, padx=4)

    def _delete_custom_var(self, index):
        if index < 0 or index >= len(self.custom_vars):
            return
        var_name = self.custom_vars[index]["name"]
        if messagebox.askyesno("确认删除", f"确定要删除变量 '{var_name}' 吗？"):
            del self.custom_vars[index]
            self._refresh_custom_var_list()
            self._refresh_all_var_buttons()

    def _refresh_custom_var_list(self):
        if not self.custom_var_container:
            return
        for w in self.custom_var_container.winfo_children():
            w.destroy()
        if not self.custom_vars:
            hint = "暂无自定义变量，请在上方添加"
            if USE_CUSTOMTK:
                ctk.CTkLabel(self.custom_var_container, text=hint, font=(FONT_FAMILY, FONT_SIZE_SMALL),
                             text_color=TEXT_SECONDARY).pack(pady=4)
            else:
                tk.Label(self.custom_var_container, text=hint, font=(FONT_FAMILY, FONT_SIZE_SMALL),
                         fg=TEXT_SECONDARY, bg=CARD).pack(pady=4)
            return
        for idx, cv in enumerate(self.custom_vars):
            row = ctk.CTkFrame(self.custom_var_container, fg_color="transparent") if USE_CUSTOMTK else tk.Frame(self.custom_var_container, bg=CARD)
            row.pack(fill=X, pady=1)
            display_text = f"{{{{{cv['name']}}}}} = {cv['value']}"
            if USE_CUSTOMTK:
                ctk.CTkLabel(row, text=display_text, font=(FONT_FAMILY, FONT_SIZE_SMALL),
                             text_color=TEXT_PRIMARY, anchor="w").pack(side=LEFT, padx=(0, 8))
                ctk.CTkButton(row, text="编辑", width=50, height=24, corner_radius=4, fg_color=BG,
                              hover_color="#D8D8D8", text_color=TEXT_SECONDARY, font=(FONT_FAMILY, FONT_SIZE_SMALL),
                              border_width=1, border_color=CARD_BORDER,
                              command=lambda i=idx: self._edit_custom_var(i)).pack(side=RIGHT, padx=2)
                ctk.CTkButton(row, text="删除", width=50, height=24, corner_radius=4, fg_color=BG,
                              hover_color="#D8D8D8", text_color=ERROR, font=(FONT_FAMILY, FONT_SIZE_SMALL),
                              border_width=1, border_color=CARD_BORDER,
                              command=lambda i=idx: self._delete_custom_var(i)).pack(side=RIGHT, padx=2)
                ctk.CTkButton(row, text="插入正文", width=64, height=24, corner_radius=4, fg_color=BG,
                              hover_color="#D8D8D8", text_color=PRIMARY, font=(FONT_FAMILY, FONT_SIZE_SMALL),
                              border_width=1, border_color=CARD_BORDER,
                              command=lambda i=idx: self.insert_var(f"{{{self.custom_vars[i]['name']}}}")).pack(side=RIGHT, padx=2)
            else:
                tk.Label(row, text=display_text, font=(FONT_FAMILY, FONT_SIZE_SMALL),
                         fg=TEXT_PRIMARY, bg=CARD, anchor="w").pack(side=LEFT, padx=(0, 8))
                tk.Button(row, text="编辑", bg=BG, fg=TEXT_SECONDARY, relief="flat", width=5, height=1,
                          font=(FONT_FAMILY, FONT_SIZE_SMALL), activebackground="#D8D8D8",
                          command=lambda i=idx: self._edit_custom_var(i)).pack(side=RIGHT, padx=2)
                tk.Button(row, text="删除", bg=BG, fg=ERROR, relief="flat", width=5, height=1,
                          font=(FONT_FAMILY, FONT_SIZE_SMALL), activebackground="#D8D8D8",
                          command=lambda i=idx: self._delete_custom_var(i)).pack(side=RIGHT, padx=2)
                tk.Button(row, text="插入正文", bg=BG, fg=PRIMARY, relief="flat", width=8, height=1,
                          font=(FONT_FAMILY, FONT_SIZE_SMALL), activebackground="#D8D8D8",
                          command=lambda i=idx: self.insert_var(f"{{{self.custom_vars[i]['name']}}}")).pack(side=RIGHT, padx=2)

    def _refresh_all_var_buttons(self):
        self._refresh_custom_var_subject_buttons()
        self._refresh_custom_var_body_buttons()

    def _refresh_custom_var_subject_buttons(self):
        if not self.custom_var_subject_container:
            return
        for w in self.custom_var_subject_container.winfo_children():
            w.destroy()
        if USE_CUSTOMTK:
            ctk.CTkLabel(self.custom_var_subject_container, text="插入变量:", font=(FONT_FAMILY, FONT_SIZE_SMALL),
                         text_color=TEXT_SECONDARY).pack(side=LEFT, padx=(0, 4))
        else:
            tk.Label(self.custom_var_subject_container, text="插入变量:", font=(FONT_FAMILY, FONT_SIZE_SMALL),
                     fg=TEXT_SECONDARY, bg=CARD).pack(side=LEFT, padx=(0, 4))
        for var in ["{姓名}", "{学号}", "{邮箱}"]:
            self._make_var_button(self.custom_var_subject_container, var, lambda v=var: self.insert_var_to_subject(v))
        for cv in self.custom_vars:
            display = f"{{{cv['name']}}}"
            self._make_var_button(self.custom_var_subject_container, display,
                                  lambda v=display: self.insert_var_to_subject(v))

    def _refresh_custom_var_body_buttons(self):
        if not self.custom_var_body_container:
            return
        for w in self.custom_var_body_container.winfo_children():
            w.destroy()
        for var in ["{姓名}", "{学号}", "{邮箱}"]:
            self._make_var_button(self.custom_var_body_container, var, lambda v=var: self.insert_var(v))
        for cv in self.custom_vars:
            display = f"{{{cv['name']}}}"
            self._make_var_button(self.custom_var_body_container, display,
                                  lambda v=display: self.insert_var(v))

    def _make_var_button(self, parent, text, cmd):
        if USE_CUSTOMTK:
            ctk.CTkButton(parent, text=text, width=max(60, len(text) * 10), height=26, corner_radius=12,
                          fg_color=BG, hover_color="#D8D8D8", text_color=PRIMARY,
                          font=(FONT_FAMILY, FONT_SIZE_SMALL), border_width=1, border_color=CARD_BORDER,
                          command=cmd).pack(side=LEFT, padx=2, pady=2)
        else:
            tk.Button(parent, text=text, bg=BG, fg=PRIMARY, relief="flat",
                      font=(FONT_FAMILY, FONT_SIZE_SMALL), activebackground="#D8D8D8",
                      command=cmd).pack(side=LEFT, padx=2, pady=2)

    def _preview_custom_vars(self):
        if not self.custom_vars:
            messagebox.showinfo("预览", "暂无自定义变量，请先添加")
            return
        dialog = tk.Toplevel(self.root)
        dialog.title("变量预览")
        dialog.geometry("500x320")
        dialog.resizable(False, False)
        dialog.transient(self.root)
        dialog.grab_set()
        if USE_CUSTOMTK:
            dialog.configure(fg_color=BG)
        else:
            dialog.configure(bg=BG)
        dlg_frame = ctk.CTkFrame(dialog, fg_color="transparent") if USE_CUSTOMTK else tk.Frame(dialog, bg=BG)
        dlg_frame.pack(fill=BOTH, expand=True, padx=20, pady=16)
        if USE_CUSTOMTK:
            ctk.CTkLabel(dlg_frame, text="变量替换预览", font=(FONT_FAMILY, FONT_SIZE_SUBTITLE, "bold"),
                         text_color=TEXT_PRIMARY).pack(anchor="w", pady=(0, 8))
        else:
            tk.Label(dlg_frame, text="变量替换预览", font=(FONT_FAMILY, FONT_SIZE_SUBTITLE, "bold"),
                     fg=TEXT_PRIMARY, bg=BG).pack(anchor="w", pady=(0, 8))
        preview_text = ctk.CTkTextbox(dlg_frame, height=200, font=(FONT_FAMILY, FONT_SIZE_BODY),
                                       fg_color=INPUT_BG, text_color=TEXT_PRIMARY) if USE_CUSTOMTK else \
                       ScrolledText(dlg_frame, height=12, font=(FONT_FAMILY, FONT_SIZE_BODY), wrap="word",
                                    bg=INPUT_BG, fg=TEXT_PRIMARY, relief="flat", padx=8, pady=8)
        preview_text.pack(fill=BOTH, expand=True, pady=(0, 8))
        lines = []
        lines.append("=== 自定义变量列表 ===\n")
        for cv in self.custom_vars:
            lines.append(f"  {{{cv['name']}}}  →  {cv['value']}")
        lines.append("")
        lines.append("=== 邮件主题预览 ===\n")
        subject = self.subject_var.get()
        for cv in self.custom_vars:
            subject = subject.replace(f"{{{cv['name']}}}", cv["value"])
        for cv_name in ["姓名", "学号", "邮箱"]:
            subject = subject.replace(f"{{{cv_name}}}", f"[{cv_name}示例]")
        lines.append(f"  {subject}")
        lines.append("")
        lines.append("=== 邮件正文预览（前200字）===\n")
        if self.body_widget:
            body = self.body_widget.get("1.0", "end-1c")
            for cv in self.custom_vars:
                body = body.replace(f"{{{cv['name']}}}", cv["value"])
            for cv_name in ["姓名", "学号", "邮箱"]:
                body = body.replace(f"{{{cv_name}}}", f"[{cv_name}示例]")
            lines.append(f"  {body[:200]}{'...' if len(body) > 200 else ''}")
        content = "\n".join(lines)
        if USE_CUSTOMTK:
            preview_text.insert("1.0", content)
            preview_text.configure(state="disabled")
        else:
            preview_text.insert("1.0", content)
            preview_text.configure(state="disabled")
        if USE_CUSTOMTK:
            ctk.CTkButton(dlg_frame, text="关闭", width=80, height=32, corner_radius=BORDER_RADIUS,
                          fg_color=PRIMARY, hover_color=PRIMARY_HOVER, text_color="white",
                          font=(FONT_FAMILY, FONT_SIZE_BODY), command=dialog.destroy).pack()
        else:
            tk.Button(dlg_frame, text="关闭", bg=PRIMARY, fg="white", relief="flat", width=8, height=1,
                      font=(FONT_FAMILY, FONT_SIZE_BODY), activebackground=PRIMARY_HOVER,
                      command=dialog.destroy).pack()

    def _build_step2_format_toolbar(self, parent):
        if USE_CUSTOMTK:
            bar = ctk.CTkFrame(parent, fg_color=CARD, corner_radius=BORDER_RADIUS, border_width=1,
                               border_color=CARD_BORDER)
            bar.pack(fill=X, padx=24, pady=6)
            row1 = ctk.CTkFrame(bar, fg_color="transparent")
            row1.pack(fill=X, padx=16, pady=(8, 2))
            ctk.CTkLabel(row1, text="格式工具", font=(FONT_FAMILY, FONT_SIZE_SMALL, "bold"),
                         text_color=TEXT_SECONDARY).pack(side=LEFT, padx=(0, 12))
            tools1 = [
                ("B", "加粗", self.format_bold),
                ("I", "斜体", self.format_italic),
                ("U", "下划线", self.format_underline),
            ]
            for sym, tip, cmd in tools1:
                ctk.CTkButton(row1, text=sym, width=32, height=28, corner_radius=6, fg_color=BG,
                              hover_color="#D8D8D8", text_color=TEXT_PRIMARY, font=(FONT_FAMILY, FONT_SIZE_BODY, "bold"),
                              border_width=1, border_color=CARD_BORDER, command=cmd).pack(side=LEFT, padx=2)
            sep1 = ctk.CTkLabel(row1, text="|", text_color=CARD_BORDER, font=(FONT_FAMILY, 14))
            sep1.pack(side=LEFT, padx=4)
            ctk.CTkLabel(row1, text="字号:", font=(FONT_FAMILY, FONT_SIZE_SMALL), text_color=TEXT_SECONDARY).pack(side=LEFT, padx=(4, 2))
            self.font_size_var = ctk.StringVar(value="正常")
            size_menu = ctk.CTkOptionMenu(row1, variable=self.font_size_var, values=["小", "正常", "大", "较大", "特大"],
                                          width=72, height=28, corner_radius=6, fg_color=BG,
                                          button_color=CARD_BORDER, button_hover_color="#D8D8D8",
                                          text_color=TEXT_PRIMARY, font=(FONT_FAMILY, FONT_SIZE_SMALL),
                                          command=self.format_font_size)
            size_menu.pack(side=LEFT, padx=2)
            sep2 = ctk.CTkLabel(row1, text="|", text_color=CARD_BORDER, font=(FONT_FAMILY, 14))
            sep2.pack(side=LEFT, padx=4)
            ctk.CTkLabel(row1, text="颜色:", font=(FONT_FAMILY, FONT_SIZE_SMALL), text_color=TEXT_SECONDARY).pack(side=LEFT, padx=(4, 2))
            self.color_var = ctk.StringVar(value="默认")
            color_menu = ctk.CTkOptionMenu(row1, variable=self.color_var,
                                           values=["默认", "红色", "蓝色", "绿色", "橙色", "紫色", "黑色"],
                                           width=72, height=28, corner_radius=6, fg_color=BG,
                                           button_color=CARD_BORDER, button_hover_color="#D8D8D8",
                                           text_color=TEXT_PRIMARY, font=(FONT_FAMILY, FONT_SIZE_SMALL),
                                           command=self.format_text_color)
            color_menu.pack(side=LEFT, padx=2)
            row2 = ctk.CTkFrame(bar, fg_color="transparent")
            row2.pack(fill=X, padx=16, pady=(2, 8))
            tools2 = [
                ("≡", "左对齐", self.format_align_left),
                ("≡", "居中对齐", self.format_align_center),
                ("≡", "右对齐", self.format_align_right),
                ("⇤", "首行缩进", self.format_first_line_indent),
                ("✕", "清除格式", self.clear_formatting),
            ]
            for sym, tip, cmd in tools2:
                ctk.CTkButton(row2, text=sym, width=32, height=28, corner_radius=6, fg_color=BG,
                              hover_color="#D8D8D8", text_color=TEXT_PRIMARY, font=(FONT_FAMILY, FONT_SIZE_BODY, "bold"),
                              border_width=1, border_color=CARD_BORDER, command=cmd).pack(side=LEFT, padx=2)
        else:
            bar = tk.Frame(parent, bg=CARD, highlightthickness=1, highlightbackground=CARD_BORDER)
            bar.pack(fill=X, padx=24, pady=6)
            row1 = tk.Frame(bar, bg=CARD)
            row1.pack(fill=X, padx=16, pady=(8, 2))
            tk.Label(row1, text="格式工具", font=(FONT_FAMILY, FONT_SIZE_SMALL, "bold"),
                     fg=TEXT_SECONDARY, bg=CARD).pack(side=LEFT, padx=(0, 12))
            for sym, tip, cmd in [("B", "加粗", self.format_bold), ("I", "斜体", self.format_italic),
                                   ("U", "下划线", self.format_underline)]:
                tk.Button(row1, text=sym, bg=BG, fg=TEXT_PRIMARY, relief="flat", width=3, height=1,
                          font=(FONT_FAMILY, FONT_SIZE_BODY, "bold"), activebackground="#D8D8D8",
                          command=cmd).pack(side=LEFT, padx=2)
            tk.Label(row1, text="|", fg=CARD_BORDER, bg=CARD, font=(FONT_FAMILY, 14)).pack(side=LEFT, padx=4)
            tk.Label(row1, text="字号:", font=(FONT_FAMILY, FONT_SIZE_SMALL), fg=TEXT_SECONDARY, bg=CARD).pack(side=LEFT, padx=(4, 2))
            self.font_size_var = tk.StringVar(value="正常")
            size_values = ["小", "正常", "大", "较大", "特大"]
            size_menu = ttk.Combobox(row1, textvariable=self.font_size_var, values=size_values, width=6, state="readonly",
                                     font=(FONT_FAMILY, FONT_SIZE_SMALL))
            size_menu.pack(side=LEFT, padx=2)
            size_menu.bind("<<ComboboxSelected>>", lambda e: self.format_font_size(self.font_size_var.get()))
            tk.Label(row1, text="|", fg=CARD_BORDER, bg=CARD, font=(FONT_FAMILY, 14)).pack(side=LEFT, padx=4)
            tk.Label(row1, text="颜色:", font=(FONT_FAMILY, FONT_SIZE_SMALL), fg=TEXT_SECONDARY, bg=CARD).pack(side=LEFT, padx=(4, 2))
            self.color_var = tk.StringVar(value="默认")
            color_values = ["默认", "红色", "蓝色", "绿色", "橙色", "紫色", "黑色"]
            color_menu = ttk.Combobox(row1, textvariable=self.color_var, values=color_values, width=6, state="readonly",
                                      font=(FONT_FAMILY, FONT_SIZE_SMALL))
            color_menu.pack(side=LEFT, padx=2)
            color_menu.bind("<<ComboboxSelected>>", lambda e: self.format_text_color(self.color_var.get()))
            row2 = tk.Frame(bar, bg=CARD)
            row2.pack(fill=X, padx=16, pady=(2, 8))
            for sym, tip, cmd in [("≡", "左对齐", self.format_align_left), ("≡", "居中对齐", self.format_align_center),
                                   ("≡", "右对齐", self.format_align_right), ("⇤", "首行缩进", self.format_first_line_indent),
                                   ("✕", "清除格式", self.clear_formatting)]:
                tk.Button(row2, text=sym, bg=BG, fg=TEXT_PRIMARY, relief="flat", width=3, height=1,
                          font=(FONT_FAMILY, FONT_SIZE_BODY, "bold"), activebackground="#D8D8D8",
                          command=cmd).pack(side=LEFT, padx=2)

    def create_step3(self):
        if USE_CUSTOMTK:
            frame = ctk.CTkFrame(self.content_frame, fg_color=BG)
        else:
            frame = tk.Frame(self.content_frame, bg=BG)
        self._card(frame, "SMTP 服务器", self._build_step3_smtp)
        self._card(frame, "发件人信息", self._build_step3_sender)
        self._build_step3_test_btn(frame)
        self._card(frame, "发送设置", self._build_step3_settings)
        return frame

    def _build_step3_test_btn(self, parent):
        if USE_CUSTOMTK:
            row = ctk.CTkFrame(parent, fg_color="transparent")
            row.pack(fill=X, padx=24, pady=6)
            ctk.CTkButton(row, text="测试连接", width=120, height=36, corner_radius=BORDER_RADIUS,
                          fg_color=SUCCESS, hover_color=SUCCESS_HOVER, text_color="white",
                          font=(FONT_FAMILY, FONT_SIZE_BUTTON, "bold"),
                          command=self.test_connection).pack(side=LEFT)
        else:
            row = tk.Frame(parent, bg=BG)
            row.pack(fill=X, padx=24, pady=6)
            tk.Button(row, text="测试连接", bg=SUCCESS, fg="white", relief="flat",
                      font=(FONT_FAMILY, FONT_SIZE_BUTTON, "bold"),
                      activebackground=SUCCESS_HOVER, command=self.test_connection).pack(side=LEFT)

    def _build_step3_smtp(self, parent):
        grid = ctk.CTkFrame(parent, fg_color="transparent") if USE_CUSTOMTK else tk.Frame(parent, bg=CARD)
        grid.pack(fill=X)
        configs = [("SMTP 服务器:", self.smtp_server, 240), ("端口:", self.smtp_port, 100)]
        for i, (label, var, w) in enumerate(configs):
            r = ctk.CTkFrame(grid, fg_color="transparent") if USE_CUSTOMTK else tk.Frame(grid, bg=CARD)
            r.pack(fill=X, pady=2)
            if USE_CUSTOMTK:
                ctk.CTkLabel(r, text=label, width=120, anchor="w", font=(FONT_FAMILY, FONT_SIZE_BODY),
                             text_color=TEXT_PRIMARY).pack(side=LEFT)
                ctk.CTkEntry(r, textvariable=var, width=w, height=32, border_width=1,
                             border_color=INPUT_BORDER, fg_color=INPUT_BG, text_color=TEXT_PRIMARY,
                             font=(FONT_FAMILY, FONT_SIZE_BODY)).pack(side=LEFT, padx=4)
            else:
                tk.Label(r, text=label, width=16, anchor="w", font=(FONT_FAMILY, FONT_SIZE_BODY),
                         fg=TEXT_PRIMARY, bg=CARD).pack(side=LEFT)
                tk.Entry(r, textvariable=var, width=w // 8, font=(FONT_FAMILY, FONT_SIZE_BODY), relief="flat",
                         bg=INPUT_BG, fg=TEXT_PRIMARY, highlightthickness=1, highlightcolor=INPUT_FOCUS,
                         highlightbackground=INPUT_BORDER).pack(side=LEFT, padx=4)
        row_tls = ctk.CTkFrame(grid, fg_color="transparent") if USE_CUSTOMTK else tk.Frame(grid, bg=CARD)
        row_tls.pack(fill=X, pady=2)
        if USE_CUSTOMTK:
            ctk.CTkLabel(row_tls, text="", width=120).pack(side=LEFT)
            ctk.CTkCheckBox(row_tls, text="使用 TLS 加密", variable=self.use_tls, onvalue=True, offvalue=False,
                            font=(FONT_FAMILY, FONT_SIZE_BODY), text_color=TEXT_PRIMARY,
                            fg_color=PRIMARY, hover_color=PRIMARY_HOVER).pack(side=LEFT)
        else:
            tk.Checkbutton(row_tls, text="使用 TLS 加密", variable=self.use_tls, font=(FONT_FAMILY, FONT_SIZE_BODY),
                           fg=TEXT_PRIMARY, bg=CARD, selectcolor=CARD, activebackground=CARD).pack(side=LEFT, padx=(120, 0))

    def _build_step3_sender(self, parent):
        configs = [("发件人邮箱:", self.sender_email), ("授权码:", self.sender_password)]
        for label, var in configs:
            row = ctk.CTkFrame(parent, fg_color="transparent") if USE_CUSTOMTK else tk.Frame(parent, bg=CARD)
            row.pack(fill=X, pady=2)
            if USE_CUSTOMTK:
                ctk.CTkLabel(row, text=label, width=120, anchor="w", font=(FONT_FAMILY, FONT_SIZE_BODY),
                             text_color=TEXT_PRIMARY).pack(side=LEFT)
                show = label.startswith("授权码")
                entry = ctk.CTkEntry(row, textvariable=var, width=300, height=32, border_width=1,
                                     border_color=INPUT_BORDER, fg_color=INPUT_BG, text_color=TEXT_PRIMARY,
                                     font=(FONT_FAMILY, FONT_SIZE_BODY), show="●" if show else "")
            else:
                tk.Label(row, text=label, width=16, anchor="w", font=(FONT_FAMILY, FONT_SIZE_BODY),
                         fg=TEXT_PRIMARY, bg=CARD).pack(side=LEFT)
                entry = tk.Entry(row, textvariable=var, width=38, font=(FONT_FAMILY, FONT_SIZE_BODY), relief="flat",
                                 bg=INPUT_BG, fg=TEXT_PRIMARY, show="●" if "授权码" in label else "",
                                 highlightthickness=1, highlightcolor=INPUT_FOCUS, highlightbackground=INPUT_BORDER)
            entry.pack(side=LEFT, padx=4)

    def _build_step3_settings(self, parent):
        settings = [
            ("发送间隔 (秒):", self.send_interval, 100),
            ("最大重试次数:", self.max_retries, 100),
            ("重试间隔 (秒):", self.retry_interval, 100),
        ]
        for label, var, w in settings:
            row = ctk.CTkFrame(parent, fg_color="transparent") if USE_CUSTOMTK else tk.Frame(parent, bg=CARD)
            row.pack(fill=X, pady=2)
            if USE_CUSTOMTK:
                ctk.CTkLabel(row, text=label, width=140, anchor="w", font=(FONT_FAMILY, FONT_SIZE_BODY),
                             text_color=TEXT_PRIMARY).pack(side=LEFT)
                ctk.CTkEntry(row, textvariable=var, width=w, height=32, border_width=1,
                             border_color=INPUT_BORDER, fg_color=INPUT_BG, text_color=TEXT_PRIMARY,
                             font=(FONT_FAMILY, FONT_SIZE_BODY)).pack(side=LEFT, padx=4)
            else:
                tk.Label(row, text=label, width=20, anchor="w", font=(FONT_FAMILY, FONT_SIZE_BODY),
                         fg=TEXT_PRIMARY, bg=CARD).pack(side=LEFT)
                tk.Entry(row, textvariable=var, width=w // 8, font=(FONT_FAMILY, FONT_SIZE_BODY), relief="flat",
                         bg=INPUT_BG, fg=TEXT_PRIMARY, highlightthickness=1, highlightcolor=INPUT_FOCUS,
                         highlightbackground=INPUT_BORDER).pack(side=LEFT, padx=4)

    def create_step4(self):
        if USE_CUSTOMTK:
            frame = ctk.CTkFrame(self.content_frame, fg_color=BG)
        else:
            frame = tk.Frame(self.content_frame, bg=BG)
        self._card(frame, "发送进度", self._build_step4_progress)
        self._card(frame, "发送日志", self._build_step4_log)
        self._build_step4_controls(frame)
        return frame

    def _build_step4_progress(self, parent):
        if USE_CUSTOMTK:
            self.progress_label = ctk.CTkLabel(parent, text="等待发送…", font=(FONT_FAMILY, FONT_SIZE_BODY),
                                               text_color=TEXT_SECONDARY, anchor="w")
            self.progress_label.pack(fill=X, pady=(0, 4))
            self.progress_bar = ctk.CTkProgressBar(parent, width=400, height=16, corner_radius=8,
                                                    fg_color=BG, progress_color=PRIMARY, border_width=0)
        else:
            self.progress_label = tk.Label(parent, text="等待发送…", font=(FONT_FAMILY, FONT_SIZE_BODY),
                                           fg=TEXT_SECONDARY, bg=CARD, anchor="w")
            self.progress_label.pack(fill=X, pady=(0, 4))
            self.progress_bar = ttk.Progressbar(parent, orient=HORIZONTAL, length=400, mode="determinate")
        self.progress_bar.pack(fill=X, pady=(0, 4))
        if USE_CUSTOMTK:
            self.progress_bar.set(0)

    def _build_step4_log(self, parent):
        self.log_widget = ScrolledText(parent, height=14, font=("Consolas", FONT_SIZE_SMALL), wrap=WORD,
                                       bg="#1E1E1E", fg="#D4D4D4", relief="flat", insertbackground="#D4D4D4",
                                       padx=8, pady=8)
        self.log_widget.pack(fill=BOTH, expand=True, pady=(0, 4))
        self.log_widget.configure(state="disabled")

    def _build_step4_controls(self, parent):
        if USE_CUSTOMTK:
            bar = ctk.CTkFrame(parent, fg_color="transparent")
            bar.pack(fill=X, padx=24, pady=6)
            self.pause_btn = ctk.CTkButton(bar, text="暂停", width=80, height=36, corner_radius=BORDER_RADIUS,
                                           fg_color=WARNING, hover_color="#E07B00", text_color="white",
                                           font=(FONT_FAMILY, FONT_SIZE_BUTTON), command=self.toggle_pause, state="disabled")
            self.pause_btn.pack(side=LEFT, padx=4)
            self.stop_btn = ctk.CTkButton(bar, text="停止", width=80, height=36, corner_radius=BORDER_RADIUS,
                                          fg_color=ERROR, hover_color="#B02A30", text_color="white",
                                          font=(FONT_FAMILY, FONT_SIZE_BUTTON), command=self.stop_send, state="disabled")
            self.stop_btn.pack(side=LEFT, padx=4)
            ctk.CTkButton(bar, text="导出失败记录", width=120, height=36, corner_radius=BORDER_RADIUS,
                          fg_color=TEXT_SECONDARY, hover_color="#4A4A4A", text_color="white",
                          font=(FONT_FAMILY, FONT_SIZE_BUTTON), command=self.export_failed).pack(side=RIGHT, padx=4)
        else:
            bar = tk.Frame(parent, bg=BG)
            bar.pack(fill=X, padx=24, pady=6)
            self.pause_btn = tk.Button(bar, text="暂停", bg=WARNING, fg="white", relief="flat",
                                       font=(FONT_FAMILY, FONT_SIZE_BUTTON), activebackground="#E07B00",
                                       command=self.toggle_pause, state="disabled")
            self.pause_btn.pack(side=LEFT, padx=4)
            self.stop_btn = tk.Button(bar, text="停止", bg=ERROR, fg="white", relief="flat",
                                      font=(FONT_FAMILY, FONT_SIZE_BUTTON), activebackground="#B02A30",
                                      command=self.stop_send, state="disabled")
            self.stop_btn.pack(side=LEFT, padx=4)
            tk.Button(bar, text="导出失败记录", bg=TEXT_SECONDARY, fg="white", relief="flat",
                      font=(FONT_FAMILY, FONT_SIZE_BUTTON), activebackground="#4A4A4A",
                      command=self.export_failed).pack(side=RIGHT, padx=4)

    # --------------------------------------------------- navigation
    def show_step(self, index):
        for f in self.step_frames:
            f.pack_forget()
        self.step_frames[index].pack(fill=BOTH, expand=True)
        self.current_step = index
        self._update_step_buttons()
        self._update_nav_buttons()

    def _update_nav_buttons(self):
        is_first = self.current_step == 0
        is_last = self.current_step == len(self.step_frames) - 1
        if USE_CUSTOMTK:
            self.prev_btn.configure(state="disabled" if is_first else "normal")
            if is_last:
                self.next_btn.pack_forget()
                self.send_btn.pack(side=RIGHT)
            else:
                self.send_btn.pack_forget()
                self.next_btn.pack(side=RIGHT)
        else:
            self.prev_btn.configure(state="disabled" if is_first else "normal")
            if is_last:
                self.next_btn.pack_forget()
                self.send_btn.pack(side=RIGHT)
            else:
                self.send_btn.pack_forget()
                self.next_btn.pack(side=RIGHT)

    def go_to_step(self, index):
        if index < 0 or index >= len(self.step_frames):
            return
        if index > self.current_step + 1:
            return
        if index > self.current_step:
            for s in range(self.current_step, index):
                if not self.validate_step(s):
                    return
        self.show_step(index)

    def prev_step(self):
        if self.current_step > 0:
            self.show_step(self.current_step - 1)

    def next_step(self):
        if self.current_step < len(self.step_frames) - 1:
            if self.validate_step(self.current_step):
                self.show_step(self.current_step + 1)

    def validate_step(self, index):
        if index == 0:
            if not self.folder_path.get().strip():
                messagebox.showwarning("提示", "请选择附件文件夹路径")
                return False
            if not os.path.isdir(self.folder_path.get().strip()):
                messagebox.showerror("错误", "附件文件夹路径不存在")
                return False
            if not self.excel_path.get().strip():
                messagebox.showwarning("提示", "请选择 Excel 文件路径")
                return False
            if not os.path.isfile(self.excel_path.get().strip()):
                messagebox.showerror("错误", "Excel 文件不存在")
                return False
            if self.column_mapping.get_column_index("email") < 0:
                messagebox.showwarning("提示", "请配置邮箱列映射")
                return False
            return True
        if index == 1:
            if not self.subject_var.get().strip():
                messagebox.showwarning("提示", "请输入邮件主题")
                return False
            body = self.body_widget.get("1.0", "end-1c")
            if not body.strip():
                messagebox.showwarning("提示", "请输入邮件正文")
                return False
            return True
        if index == 2:
            if not self.smtp_server.get().strip():
                messagebox.showwarning("提示", "请输入 SMTP 服务器地址")
                return False
            try:
                port = self.smtp_port.get()
                if not (1 <= port <= 65535):
                    messagebox.showwarning("提示", "端口号必须在 1-65535 之间")
                    return False
            except Exception:
                messagebox.showwarning("提示", "端口号必须是有效的数字")
                return False
            if not self.sender_email.get().strip():
                messagebox.showwarning("提示", "请输入发件人邮箱")
                return False
            if not self.sender_password.get().strip():
                messagebox.showwarning("提示", "请输入授权码")
                return False
            if not self.validate_interval() or not self.validate_retries():
                return False
            return True
        return True

    # ------------------------------------------------- step1 actions
    def browse_folder(self):
        path = filedialog.askdirectory(title="选择附件文件夹")
        if path:
            self.folder_path.set(path)

    def browse_excel(self):
        path = filedialog.askopenfilename(title="选择 Excel 文件", filetypes=[("Excel 文件", "*.xlsx *.xls")])
        if path:
            self.excel_path.set(path)
            self.load_columns()

    def load_columns(self):
        path = self.excel_path.get().strip()
        if not path or not os.path.isfile(path):
            messagebox.showerror("错误", "请先选择有效的 Excel 文件")
            return
        try:
            self.column_mapping.load_from_excel(path)
            self._build_mapping_widgets()
            self.load_preview()
        except Exception as e:
            messagebox.showerror("错误", f"加载 Excel 失败: {e}")

    def load_preview(self):
        for w in self.preview_container.winfo_children():
            w.destroy()
        folder = self.folder_path.get().strip()
        if not folder or not os.path.isdir(folder):
            if self.preview_hint:
                self.preview_hint.configure(text="请先选择有效的附件文件夹")
            return
        if not self.column_mapping.headers:
            if self.preview_hint:
                self.preview_hint.configure(text="请先加载 Excel 文件并配置列映射")
            return
        if not self.column_mapping.mapping:
            if self.preview_hint:
                self.preview_hint.configure(text="请先配置关键列映射")
            return
        key_field_name = self.column_mapping.get_key_field_name()
        matcher = RecipientMatcher(folder, self.column_mapping, key_field_name)
        self.match_results = matcher.match_files()
        matched = [r for r in self.match_results if r["matched"]]
        not_matched = [r for r in self.match_results if not r["matched"]]
        if self.preview_hint:
            if matched:
                self.preview_hint.configure(
                    text=f"共 {len(self.match_results)} 个文件，匹配成功 {len(matched)} 个，未匹配 {len(not_matched)} 个",
                    text_color=SUCCESS if USE_CUSTOMTK else SUCCESS)
            else:
                self.preview_hint.configure(
                    text=f"共 {len(self.match_results)} 个文件，无匹配收件人！请检查列映射和文件名",
                    text_color=WARNING if USE_CUSTOMTK else WARNING)
        headers = ["文件名", "关键值", "收件人", "状态"]
        if USE_CUSTOMTK:
            header_row = ctk.CTkFrame(self.preview_container, fg_color=PRIMARY)
            header_row.pack(fill=X, pady=(0, 2))
            for h in headers:
                ctk.CTkLabel(header_row, text=h, font=(FONT_FAMILY, FONT_SIZE_SMALL, "bold"),
                             text_color="white", width=120, anchor="w").pack(side=LEFT, padx=4, pady=2)
        else:
            header_row = tk.Frame(self.preview_container, bg=PRIMARY)
            header_row.pack(fill=X, pady=(0, 2))
            for h in headers:
                tk.Label(header_row, text=h, font=(FONT_FAMILY, FONT_SIZE_SMALL, "bold"),
                         fg="white", bg=PRIMARY, width=16, anchor="w").pack(side=LEFT, padx=4, pady=2)
        email_idx = self.column_mapping.get_column_index("email")
        name_idx = self.column_mapping.get_column_index("name")
        for r in self.match_results:
            row_bg = CARD
            if USE_CUSTOMTK:
                row_frame = ctk.CTkFrame(self.preview_container, fg_color=row_bg, corner_radius=4)
            else:
                row_frame = tk.Frame(self.preview_container, bg=row_bg)
            row_frame.pack(fill=X, pady=1)
            ctk_or_tk_label = ctk.CTkLabel if USE_CUSTOMTK else tk.Label
            label_kwargs = {"font": (FONT_FAMILY, FONT_SIZE_SMALL)}
            if not USE_CUSTOMTK:
                label_kwargs["bg"] = row_bg
            ctk_or_tk_label(row_frame, text=r["filename"], width=160, anchor="w",
                            text_color=TEXT_PRIMARY, **label_kwargs).pack(side=LEFT, padx=4)
            ctk_or_tk_label(row_frame, text=r.get("key", ""), width=80, anchor="w",
                            text_color=TEXT_SECONDARY, **label_kwargs).pack(side=LEFT, padx=4)
            if r["matched"]:
                row_idx = r["row_idx"]
                name_val = ""
                email_val = ""
                if row_idx is not None and name_idx >= 0:
                    data = self.column_mapping.data
                    if row_idx < len(data) and name_idx < len(data[row_idx]):
                        name_val = str(data[row_idx][name_idx] or "")
                if row_idx is not None and email_idx >= 0:
                    data = self.column_mapping.data
                    if row_idx < len(data) and email_idx < len(data[row_idx]):
                        email_val = str(data[row_idx][email_idx] or "")
                recipient_text = f"{name_val} <{email_val}>" if email_val else name_val
                ctk_or_tk_label(row_frame, text=recipient_text, width=200, anchor="w",
                                text_color=TEXT_PRIMARY, **label_kwargs).pack(side=LEFT, padx=4)
                status_color = SUCCESS
                status_text = "匹配成功"
            else:
                ctk_or_tk_label(row_frame, text="", width=200, anchor="w",
                                text_color=TEXT_SECONDARY, **label_kwargs).pack(side=LEFT, padx=4)
                status_color = WARNING
                status_text = "未匹配"
            if USE_CUSTOMTK:
                ctk_or_tk_label(row_frame, text=status_text, width=80, anchor="w",
                                text_color=status_color, **label_kwargs).pack(side=LEFT, padx=4)
            else:
                tk.Label(row_frame, text=status_text, width=12, anchor="w",
                         fg=status_color, bg=row_bg, font=(FONT_FAMILY, FONT_SIZE_SMALL)).pack(side=LEFT, padx=4)

    # ------------------------------------------------- step2 actions
    def insert_var(self, var):
        if self.body_widget:
            self.body_widget.insert("insert", var)

    def insert_var_to_subject(self, var):
        current = self.subject_var.get()
        self.subject_var.set(current + var)

    def _get_sel(self):
        try:
            sel = self.body_widget.tag_ranges("sel")
            return (sel[0], sel[1]) if sel else None
        except Exception:
            return None

    def _get_font_style_state(self, tags):
        font_str = "normal"
        underline = False
        try:
            if "font_style" in tags:
                font_cfg = self.body_widget.tag_cget("font_style", "font")
                if font_cfg:
                    parts = str(font_cfg).split()
                    if len(parts) >= 3:
                        font_str = parts[2] if len(parts) > 2 else "normal"
                underline_cfg = self.body_widget.tag_cget("font_style", "underline")
                underline = str(underline_cfg) == "1"
        except Exception:
            pass
        styles = set(font_str.split()) if font_str != "normal" else set()
        return {"bold": "bold" in styles, "italic": "italic" in styles, "underline": underline}

    def _apply_combined_font_tag(self, target_tag):
        if not self.body_widget:
            return
        try:
            sel = self._get_sel()
            if not sel:
                return
            start, end = sel
            sample_pos = start
            current_tags = self.body_widget.tag_names(sample_pos)
            state = self._get_font_style_state(current_tags)
            if target_tag == "bold":
                state["bold"] = not state["bold"]
            elif target_tag == "italic":
                state["italic"] = not state["italic"]
            elif target_tag == "underline":
                state["underline"] = not state["underline"]
            self.body_widget.tag_remove("font_style", start, end)
            combo_parts = []
            if state["bold"]:
                combo_parts.append("bold")
            if state["italic"]:
                combo_parts.append("italic")
            font_style = " ".join(combo_parts) if combo_parts else "normal"
            self.body_widget.tag_configure("font_style", font=(FONT_FAMILY, FONT_SIZE_BODY, font_style),
                                          underline=state["underline"])
            self.body_widget.tag_add("font_style", start, end)
        except Exception:
            pass

    def format_bold(self):
        self._apply_combined_font_tag("bold")

    def format_italic(self):
        self._apply_combined_font_tag("italic")

    def format_underline(self):
        self._apply_combined_font_tag("underline")

    def format_align_left(self):
        if not self.body_widget:
            return
        try:
            sel = self._get_sel()
            if sel:
                start, end = sel
                self.body_widget.tag_remove("align_right", start, end)
                self.body_widget.tag_remove("align_center", start, end)
                self.body_widget.tag_add("align_left", start, end)
            else:
                line = self.body_widget.index("insert").split(".")[0]
                line_start = f"{line}.0"
                line_end = f"{line}.end"
                self.body_widget.tag_remove("align_right", line_start, line_end)
                self.body_widget.tag_remove("align_center", line_start, line_end)
                self.body_widget.tag_add("align_left", line_start, line_end)
        except Exception:
            pass

    def format_align_center(self):
        if not self.body_widget:
            return
        try:
            sel = self._get_sel()
            if sel:
                start, end = sel
                self.body_widget.tag_remove("align_left", start, end)
                self.body_widget.tag_remove("align_right", start, end)
                self.body_widget.tag_add("align_center", start, end)
            else:
                line = self.body_widget.index("insert").split(".")[0]
                line_start = f"{line}.0"
                line_end = f"{line}.end"
                self.body_widget.tag_remove("align_left", line_start, line_end)
                self.body_widget.tag_remove("align_right", line_start, line_end)
                self.body_widget.tag_add("align_center", line_start, line_end)
        except Exception:
            pass

    def format_align_right(self):
        if not self.body_widget:
            return
        try:
            sel = self._get_sel()
            if sel:
                start, end = sel
                self.body_widget.tag_remove("align_left", start, end)
                self.body_widget.tag_remove("align_center", start, end)
                self.body_widget.tag_add("align_right", start, end)
            else:
                line = self.body_widget.index("insert").split(".")[0]
                line_start = f"{line}.0"
                line_end = f"{line}.end"
                self.body_widget.tag_remove("align_left", line_start, line_end)
                self.body_widget.tag_remove("align_center", line_start, line_end)
                self.body_widget.tag_add("align_right", line_start, line_end)
        except Exception:
            pass

    def format_font_size(self, size_name):
        if not self.body_widget:
            return
        size_map = {"小": "size_small", "正常": "size_normal", "大": "size_large", "较大": "size_xlarge", "特大": "size_xxlarge"}
        tag = size_map.get(size_name, "size_normal")
        try:
            sel = self._get_sel()
            if sel:
                start, end = sel
                for t in ["size_small", "size_normal", "size_large", "size_xlarge", "size_xxlarge"]:
                    self.body_widget.tag_remove(t, start, end)
                self.body_widget.tag_add(tag, start, end)
        except Exception:
            pass

    def format_text_color(self, color_name):
        if not self.body_widget:
            return
        color_map = {"默认": TEXT_PRIMARY, "红色": "#D13438", "蓝色": "#0078D4", "绿色": "#107C10", "橙色": "#FF8C00", "紫色": "#8764B8", "黑色": "#1A1A1A"}
        color = color_map.get(color_name, TEXT_PRIMARY)
        try:
            sel = self._get_sel()
            if sel:
                start, end = sel
                for t in list(self.body_widget.tag_names()):
                    if t.startswith("fgcolor_"):
                        self.body_widget.tag_remove(t, start, end)
                fg_tag = f"fgcolor_{color}"
                try:
                    self.body_widget.tag_cget(fg_tag, "foreground")
                except Exception:
                    self.body_widget.tag_configure(fg_tag, foreground=color)
                self.body_widget.tag_add(fg_tag, start, end)
        except Exception:
            pass

    def format_first_line_indent(self):
        if not self.body_widget:
            return
        try:
            sel = self._get_sel()
            if sel:
                start, end = sel
                current_tags = self.body_widget.tag_names(start)
                if "indented" in current_tags:
                    self.body_widget.tag_remove("indented", start, end)
                else:
                    self.body_widget.tag_add("indented", start, end)
            else:
                line = self.body_widget.index("insert").split(".")[0]
                line_start = f"{line}.0"
                line_end = f"{line}.end"
                current_tags = self.body_widget.tag_names(line_start)
                if "indented" in current_tags:
                    self.body_widget.tag_remove("indented", line_start, line_end)
                else:
                    self.body_widget.tag_add("indented", line_start, line_end)
        except Exception:
            pass

    def clear_formatting(self):
        if not self.body_widget:
            return
        try:
            tags = ["font_style", "align_left", "align_center", "align_right", "indented",
                    "size_small", "size_normal", "size_large", "size_xlarge", "size_xxlarge"]
            sel = self._get_sel()
            if sel:
                start, end = sel
                for tag in tags:
                    self.body_widget.tag_remove(tag, start, end)
                for t in list(self.body_widget.tag_names()):
                    if t.startswith("fgcolor_"):
                        self.body_widget.tag_remove(t, start, end)
            else:
                for tag in tags:
                    self.body_widget.tag_remove(tag, "1.0", "end")
                for t in list(self.body_widget.tag_names()):
                    if t.startswith("fgcolor_"):
                        self.body_widget.tag_remove(t, "1.0", "end")
        except Exception:
            pass

    def get_formatted_body_html(self):
        text = self.body_widget.get("1.0", "end-1c")
        lines = text.split("\n")
        html_parts = []
        for i, line in enumerate(lines):
            if not line.strip():
                html_parts.append("<br>")
                continue
            line_num = i + 1
            chars = list(line)
            align_right_positions = set()
            align_center_positions = set()
            indent_positions = set()
            bold_positions = set()
            italic_positions = set()
            underline_positions = set()
            size_positions = {}
            color_positions = {}
            size_tags = ["size_small", "size_normal", "size_large", "size_xlarge", "size_xxlarge"]
            for j in range(len(chars)):
                pos = f"{line_num}.{j}"
                try:
                    tags = self.body_widget.tag_names(pos)
                    if "align_right" in tags:
                        align_right_positions.add(j)
                    if "align_center" in tags:
                        align_center_positions.add(j)
                    if "indented" in tags:
                        indent_positions.add(j)
                    if "font_style" in tags:
                        try:
                            font_cfg = str(self.body_widget.tag_cget("font_style", "font"))
                            parts = font_cfg.split()
                            font_style_str = parts[2] if len(parts) >= 3 else "normal"
                            style_set = set(font_style_str.split()) if font_style_str != "normal" else set()
                            if "bold" in style_set:
                                bold_positions.add(j)
                            if "italic" in style_set:
                                italic_positions.add(j)
                            underline_cfg = self.body_widget.tag_cget("font_style", "underline")
                            if str(underline_cfg) == "1":
                                underline_positions.add(j)
                        except Exception:
                            pass
                    for st in size_tags:
                        if st in tags:
                            size_positions[j] = st
                            break
                    for t in tags:
                        if t.startswith("fgcolor_"):
                            color_positions[j] = t.replace("fgcolor_", "")
                            break
                except Exception:
                    pass
            has_right = len(align_right_positions) > len(chars) // 2 if chars else False
            has_center = len(align_center_positions) > len(chars) // 2 if chars else False
            has_indent = len(indent_positions) > len(chars) // 2 if chars else False
            result = ""
            current_state = {"bold": False, "italic": False, "underline": False}
            for j, char in enumerate(chars):
                want_bold = j in bold_positions
                want_italic = j in italic_positions
                want_underline = j in underline_positions
                if want_bold != current_state["bold"] or want_italic != current_state["italic"] or want_underline != current_state["underline"]:
                    if any(current_state[k] for k in ("bold", "italic", "underline")):
                        result += "</span>"
                    current_state["bold"] = want_bold
                    current_state["italic"] = want_italic
                    current_state["underline"] = want_underline
                    if any(current_state[k] for k in ("bold", "italic", "underline")):
                        style_parts = []
                        if current_state["bold"]:
                            style_parts.append("font-weight:bold")
                        if current_state["italic"]:
                            style_parts.append("font-style:italic")
                        if current_state["underline"]:
                            style_parts.append("text-decoration:underline")
                        result += f'<span style="{";".join(style_parts)}">'
                escaped = char.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
                result += escaped
            if any(current_state[k] for k in ("bold", "italic", "underline")):
                result += "</span>"
            p_style = ""
            if has_right:
                p_style = "text-align:right;"
            elif has_center:
                p_style = "text-align:center;"
            elif has_indent:
                p_style = "text-indent:2em;"
            if p_style:
                html_parts.append(f'<p style="{p_style}">{result}</p>')
            else:
                html_parts.append(f"<p>{result}</p>")
        return "\n".join(html_parts)

    # ------------------------------------------------- step3 actions
    def test_connection(self):
        server_addr = self.smtp_server.get().strip()
        port = self.smtp_port.get()
        tls = self.use_tls.get()
        email = self.sender_email.get().strip()
        pwd = self.sender_password.get().strip()
        if not email or not pwd:
            messagebox.showwarning("提示", "请填写发件人邮箱和授权码")
            return

        def _test():
            try:
                self.root.after(0, lambda: self.log_message(f"正在测试连接 {server_addr}:{port} ..."))
                srv = smtplib.SMTP(server_addr, port, timeout=15)
                if tls:
                    srv.starttls()
                srv.login(email, pwd)
                srv.quit()
                self.root.after(0, lambda: self.log_message("SMTP 连接测试成功!"))
                self.root.after(0, lambda: messagebox.showinfo("成功", "SMTP 连接测试成功!"))
            except Exception as e:
                err = str(e)
                self.root.after(0, lambda: self.log_message(f"SMTP 连接测试失败: {err}"))
                self.root.after(0, lambda: messagebox.showerror("失败", f"SMTP 连接测试失败:\n{err}"))

        threading.Thread(target=_test, daemon=True).start()

    def validate_interval(self):
        try:
            val = self.send_interval.get()
            if val < 0:
                raise ValueError
            return True
        except Exception:
            messagebox.showwarning("提示", "发送间隔必须为非负数")
            return False

    def validate_retries(self):
        try:
            val = self.max_retries.get()
            if val < 0:
                raise ValueError
            val2 = self.retry_interval.get()
            if val2 < 0:
                raise ValueError
            return True
        except Exception:
            messagebox.showwarning("提示", "重试次数和重试间隔必须为非负整数")
            return False

    # ------------------------------------------------- step4 actions
    def start_send(self):
        if self.is_sending:
            return
        for step_idx in range(3):
            if not self.validate_step(step_idx):
                self.show_step(step_idx)
                return
        self.show_step(3)
        self.stop_event.clear()
        self.pause_event.set()
        self.is_sending = True
        self.success_count = 0
        self.fail_count = 0
        self.failed_records = []
        self._update_send_controls(True)
        self.send_thread = threading.Thread(target=self.send_emails, daemon=True)
        self.send_thread.start()

    def _update_send_controls(self, sending):
        state = "normal" if sending else "disabled"
        try:
            self.pause_btn.configure(state=state)
            self.stop_btn.configure(state=state)
        except Exception:
            pass

    def send_emails(self):
        folder = self.folder_path.get().strip()
        matcher = RecipientMatcher(folder, self.column_mapping, self.column_mapping.get_key_field_name())
        matched = matcher.match_files()
        targets = [m for m in matched if m["matched"]]
        self.total_count = len(targets)

        if self.total_count == 0:
            self.root.after(0, lambda: self.log_message("没有找到匹配的文件，请检查列映射配置"))
            self.root.after(0, self.on_send_complete)
            return

        self.root.after(0, lambda: self.log_message(f"共 {self.total_count} 个文件待发送"))

        email_idx = self.column_mapping.get_column_index("email")
        name_idx = self.column_mapping.get_column_index("name")
        subject_tpl = self.subject_var.get()
        body_html = self.get_formatted_body_html()

        interval = self.send_interval.get()
        retries = self.max_retries.get()
        retry_int = self.retry_interval.get()
        smtp_cfg = {
            "smtp_server": self.smtp_server.get().strip(),
            "smtp_port": self.smtp_port.get(),
            "use_tls": self.use_tls.get(),
            "sender_email": self.sender_email.get().strip(),
            "sender_password": self.sender_password.get().strip(),
        }

        for i, item in enumerate(targets):
            if self.stop_event.is_set():
                self.root.after(0, lambda: self.log_message("发送已停止"))
                break
            self.pause_event.wait()

            row_idx = item["row_idx"]
            data_row = self.column_mapping.data[row_idx] if row_idx is not None and row_idx < len(self.column_mapping.data) else None
            recipient_email = ""
            recipient_name = ""
            if data_row:
                if 0 <= email_idx < len(data_row) and data_row[email_idx]:
                    recipient_email = str(data_row[email_idx]).strip()
                if 0 <= name_idx < len(data_row) and data_row[name_idx]:
                    recipient_name = str(data_row[name_idx]).strip()

            if not recipient_email:
                self.fail_count += 1
                self.failed_records.append({
                    "文件名": item["filename"], "关键值": item["key"],
                    "姓名": recipient_name, "邮箱": "", "错误": "邮箱地址为空",
                    "时间": datetime.now().strftime("%Y-%m-%d %H:%M:%S")
                })
                self.root.after(0, lambda fn=item["filename"]: self.log_message(f"跳过 {fn}: 邮箱地址为空"))
                self._update_progress(i + 1)
                continue

            subject = subject_tpl
            body = body_html
            replace_map = {"{姓名}": recipient_name, "{学号}": item["key"], "{邮箱}": recipient_email}
            for cv in self.custom_vars:
                replace_map[f"{{{cv['name']}}}"] = cv["value"]
            for k, v in replace_map.items():
                subject = subject.replace(k, v)
                body = body.replace(k, v)

            success = False
            last_error = ""
            for attempt in range(retries):
                if self.stop_event.is_set():
                    break
                try:
                    self.send_single_email(smtp_cfg, recipient_email, subject, body, item["filepath"])
                    success = True
                    break
                except Exception as e:
                    last_error = str(e)
                    if attempt < retries - 1:
                        self.root.after(0, lambda err=last_error, a=attempt + 1: self.log_message(
                            f"发送失败 (第{a}次): {err}，{retry_int}秒后重试..."))
                        for _ in range(retry_int * 10):
                            if self.stop_event.is_set():
                                break
                            time.sleep(0.1)

            if success:
                self.success_count += 1
                self.root.after(0, lambda fn=item["filename"], rn=recipient_name: self.log_message(
                    f"✓ 发送成功: {fn} -> {rn}"))
            else:
                self.fail_count += 1
                self.failed_records.append({
                    "文件名": item["filename"], "关键值": item["key"],
                    "姓名": recipient_name, "邮箱": recipient_email, "错误": last_error,
                    "时间": datetime.now().strftime("%Y-%m-%d %H:%M:%S")
                })
                self.root.after(0, lambda fn=item["filename"], err=last_error: self.log_message(
                    f"✗ 发送失败: {fn} - {err}"))

            self._update_progress(i + 1)

            if i < len(targets) - 1 and not self.stop_event.is_set():
                for _ in range(int(interval * 10)):
                    if self.stop_event.is_set():
                        break
                    time.sleep(0.1)

        self.root.after(0, self.on_send_complete)

    def send_single_email(self, smtp_config, recipient, subject, body_html, attachment_path=None):
        msg = MIMEMultipart("mixed")
        msg["From"] = smtp_config["sender_email"]
        msg["To"] = recipient
        msg["Subject"] = Header(subject, "utf-8").encode()

        alt_part = MIMEMultipart("alternative")
        plain_text = re.sub(r"<[^>]+>", "", body_html)
        alt_part.attach(MIMEText(plain_text, "plain", "utf-8"))
        alt_part.attach(MIMEText(body_html, "html", "utf-8"))
        msg.attach(alt_part)

        if attachment_path and os.path.isfile(attachment_path):
            with open(attachment_path, "rb") as f:
                part = MIMEApplication(f.read(), Name=os.path.basename(attachment_path))
            part.add_header("Content-Disposition", "attachment", filename=("utf-8", "", os.path.basename(attachment_path)))
            msg.attach(part)

        server = None
        try:
            port = smtp_config.get("smtp_port", 587)
            timeout_val = 30
            if port == 465:
                server = smtplib.SMTP_SSL(smtp_config["smtp_server"], port, timeout=timeout_val)
            else:
                server = smtplib.SMTP(smtp_config["smtp_server"], port, timeout=timeout_val)
                if smtp_config.get("use_tls", True):
                    server.starttls()
            server.login(smtp_config["sender_email"], smtp_config["sender_password"])
            server.sendmail(smtp_config["sender_email"], [recipient], msg.as_string())
        finally:
            if server:
                try:
                    server.quit()
                except Exception:
                    pass

    def _update_progress(self, current):
        total = max(self.total_count, 1)
        pct = current / total
        self.root.after(0, lambda: self._set_progress(pct, current))

    def _set_progress(self, pct, current):
        if USE_CUSTOMTK:
            self.progress_bar.set(pct)
        else:
            self.progress_bar["value"] = pct * 100
        self.progress_label.configure(text=f"进度: {current}/{self.total_count}  |  成功: {self.success_count}  |  失败: {self.fail_count}")

    def toggle_pause(self):
        if not self.is_sending:
            return
        if self.pause_event.is_set():
            self.pause_event.clear()
            self.root.after(0, lambda: self.log_message("发送已暂停"))
            try:
                self.pause_btn.configure(text="继续")
            except Exception:
                pass
        else:
            self.pause_event.set()
            self.root.after(0, lambda: self.log_message("发送已继续"))
            try:
                self.pause_btn.configure(text="暂停")
            except Exception:
                pass

    def stop_send(self):
        if not self.is_sending:
            return
        self.stop_event.set()
        self.pause_event.set()
        self.root.after(0, lambda: self.log_message("正在停止..."))

    def on_send_complete(self):
        self.is_sending = False
        self._update_send_controls(False)
        self.log_message(f"发送完成! 成功: {self.success_count}, 失败: {self.fail_count}")
        if self.fail_count > 0:
            messagebox.showwarning("完成", f"发送完成!\n成功: {self.success_count}\n失败: {self.fail_count}\n\n可点击「导出失败记录」保存详情")
        else:
            messagebox.showinfo("完成", f"全部发送成功!\n共 {self.success_count} 封")

    def log_message(self, message):
        if not self.log_widget:
            return
        timestamp = datetime.now().strftime("%H:%M:%S")
        line = f"[{timestamp}] {message}\n"
        try:
            self.log_widget.configure(state="normal")
            self.log_widget.insert(END, line)
            self.log_widget.see(END)
            self.log_widget.configure(state="disabled")
        except Exception:
            pass

    def export_failed(self):
        if not self.failed_records:
            messagebox.showinfo("提示", "没有失败记录可导出")
            return
        dialog = tk.Toplevel(self.root)
        dialog.title("导出失败记录")
        dialog.geometry("360x200")
        dialog.resizable(False, False)
        dialog.transient(self.root)
        dialog.grab_set()
        if USE_CUSTOMTK:
            dialog.configure(fg_color=BG)
        else:
            dialog.configure(bg=BG)
        dlg_frame = ctk.CTkFrame(dialog, fg_color="transparent") if USE_CUSTOMTK else tk.Frame(dialog, bg=BG)
        dlg_frame.pack(fill=BOTH, expand=True, padx=20, pady=16)
        if USE_CUSTOMTK:
            ctk.CTkLabel(dlg_frame, text="请选择导出格式:", font=(FONT_FAMILY, FONT_SIZE_BODY, "bold"),
                         text_color=TEXT_PRIMARY).pack(pady=(0, 12))
        else:
            tk.Label(dlg_frame, text="请选择导出格式:", font=(FONT_FAMILY, FONT_SIZE_BODY, "bold"),
                     fg=TEXT_PRIMARY, bg=BG).pack(pady=(0, 12))
        btn_frame = ctk.CTkFrame(dlg_frame, fg_color="transparent") if USE_CUSTOMTK else tk.Frame(dlg_frame, bg=BG)
        btn_frame.pack(fill=X)
        def do_export_excel():
            dialog.destroy()
            self._export_failed_excel()
        def do_export_log():
            dialog.destroy()
            self._export_failed_log()
        def do_export_both():
            dialog.destroy()
            self._export_failed_excel()
            self._export_failed_log()
        if USE_CUSTOMTK:
            ctk.CTkButton(btn_frame, text="Excel 格式", width=100, height=36, corner_radius=BORDER_RADIUS,
                          fg_color=PRIMARY, hover_color=PRIMARY_HOVER, text_color="white",
                          font=(FONT_FAMILY, FONT_SIZE_BUTTON), command=do_export_excel).pack(side=LEFT, padx=4)
            ctk.CTkButton(btn_frame, text="Log 格式", width=100, height=36, corner_radius=BORDER_RADIUS,
                          fg_color=SUCCESS, hover_color=SUCCESS_HOVER, text_color="white",
                          font=(FONT_FAMILY, FONT_SIZE_BUTTON), command=do_export_log).pack(side=LEFT, padx=4)
            ctk.CTkButton(btn_frame, text="全部导出", width=100, height=36, corner_radius=BORDER_RADIUS,
                          fg_color=WARNING, hover_color="#E07B00", text_color="white",
                          font=(FONT_FAMILY, FONT_SIZE_BUTTON, "bold"), command=do_export_both).pack(side=LEFT, padx=4)
        else:
            tk.Button(btn_frame, text="Excel 格式", bg=PRIMARY, fg="white", relief="flat", width=12, height=1,
                      font=(FONT_FAMILY, FONT_SIZE_BUTTON), activebackground=PRIMARY_HOVER,
                      command=do_export_excel).pack(side=LEFT, padx=4)
            tk.Button(btn_frame, text="Log 格式", bg=SUCCESS, fg="white", relief="flat", width=12, height=1,
                      font=(FONT_FAMILY, FONT_SIZE_BUTTON), activebackground=SUCCESS_HOVER,
                      command=do_export_log).pack(side=LEFT, padx=4)
            tk.Button(btn_frame, text="全部导出", bg=WARNING, fg="white", relief="flat", width=12, height=1,
                      font=(FONT_FAMILY, FONT_SIZE_BUTTON, "bold"), activebackground="#E07B00",
                      command=do_export_both).pack(side=LEFT, padx=4)

    def _export_failed_excel(self):
        path = filedialog.asksaveasfilename(
            title="导出失败记录 (Excel)",
            defaultextension=".xlsx",
            filetypes=[("Excel 文件", "*.xlsx")],
            initialfile=f"failed_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
        )
        if not path:
            return
        try:
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = "失败记录"
            headers = ["序号", "文件名", "关键值", "姓名", "邮箱", "错误原因", "时间"]
            ws.append(headers)
            for idx, rec in enumerate(self.failed_records, 1):
                ws.append([idx, rec.get("文件名", ""), rec.get("关键值", ""),
                           rec.get("姓名", ""), rec.get("邮箱", ""),
                           rec.get("错误", ""), rec.get("时间", "")])
            for col in range(1, len(headers) + 1):
                ws.column_dimensions[openpyxl.utils.get_column_letter(col)].width = 20
            wb.save(path)
            messagebox.showinfo("成功", f"失败记录已导出至:\n{path}")
            self.log_message(f"失败记录(Excel)已导出: {path}")
        except Exception as e:
            messagebox.showerror("错误", f"导出 Excel 失败: {e}")

    def _export_failed_log(self):
        path = filedialog.asksaveasfilename(
            title="导出失败记录 (Log)",
            defaultextension=".log",
            filetypes=[("Log 文件", "*.log"), ("文本文件", "*.txt")],
            initialfile=f"failed_{datetime.now().strftime('%Y%m%d_%H%M%S')}.log"
        )
        if not path:
            return
        try:
            with open(path, "w", encoding="utf-8") as f:
                f.write(f"{'='*70}\n")
                f.write(f"  批量邮件发送 - 失败记录\n")
                f.write(f"  导出时间: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}\n")
                f.write(f"  总失败数: {len(self.failed_records)}\n")
                f.write(f"{'='*70}\n\n")
                for idx, rec in enumerate(self.failed_records, 1):
                    f.write(f"--- 记录 {idx} ---\n")
                    f.write(f"  文件名: {rec.get('文件名', '')}\n")
                    f.write(f"  关键值: {rec.get('关键值', '')}\n")
                    f.write(f"  姓  名: {rec.get('姓名', '')}\n")
                    f.write(f"  邮  箱: {rec.get('邮箱', '')}\n")
                    f.write(f"  错  误: {rec.get('错误', '')}\n")
                    f.write(f"  时  间: {rec.get('时间', '')}\n\n")
                f.write(f"{'='*70}\n")
                f.write(f"  结束\n")
                f.write(f"{'='*70}\n")
            messagebox.showinfo("成功", f"失败记录已导出至:\n{path}")
            self.log_message(f"失败记录(Log)已导出: {path}")
        except Exception as e:
            messagebox.showerror("错误", f"导出 Log 失败: {e}")

    # ------------------------------------------------------- misc
    def on_closing(self):
        if self.is_sending:
            if messagebox.askyesno("确认退出", "正在发送邮件，确定要退出吗？"):
                self.stop_event.set()
                self.pause_event.set()
                self.root.after(500, self.root.destroy)
            return
        self.root.destroy()

    def run(self):
        self.root.mainloop()


if __name__ == "__main__":
    app = EmailSenderGUI()
    app.run()
